"""Build the Probrands database from the Professional Brands Market Map workbook.

Usage: python3 probrands/import/build_import.py <Professional_Brands_Market_Map.xlsx>

Writes import/out/data/*.json, one file per document of the app's "data" collection:
  co-N   company summaries (owner, model, SKU mix, the metrics behind the DNA test, Claude's judgment)
  dt-N   company details (write-up, brand-by-brand list, acquisitions, sources)
  br-N   every brand: owner (verified, unverified prior, Claude's guess or unknown), channels, formats
  deals  every deal in the ledger, with the acquiring and target owners resolved
  dists  the distributor list with channel class and end market
Writes import/out/sku/p-N.json, the "sku" collection: every scraped SKU, grouped by brand, loaded on demand.
Also writes import/out/judge/batch-N.json: the facts Claude needs for the judgment tests.
Judgments live in import/judgments.json (company id -> verdicts) and are merged into co-N.
Claude's best-guess owners for brands the workbook has no owner for live in import/brand_owners.json (brand name -> guess).

Owners come from the Master tab (verified ownership, with facts from Companies, Profiles and
Sources). Brands that only carry Claude's unverified "Likely owner (PRIOR)" join a Master owner
when the names match, and otherwise become owners of their own with ownership marked unverified.
"""
import collections, glob, json, math, os, re, sys, unicodedata

import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else "Professional_Brands_Market_Map.xlsx"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "out")
NOW = "2026-10-01T12:00:00Z"
VERSION = re.search(r"v\d+", os.path.basename(SRC))
SEEDED = f"Professional Brands Market Map {VERSION.group(0) if VERSION else ''} (1 Oct 2026)".replace("  ", " ")

wb = openpyxl.load_workbook(SRC, data_only=True)
FIRST_COL = {"Rank", "ID", "Brand ID", "Year", "#", "Product name", "Segment", "Distributor", "Section"}


def table(name):
    """Rows of a tab as dicts, keyed by its header row (the first row whose first cell is a known header)."""
    rows = list(wb[name].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] in FIRST_COL)
    head, seen = [], collections.Counter()
    for h in rows[hi]:  # duplicate headers ("% Other" twice) get a suffix
        h = str(h).strip() if h is not None else ""
        seen[h] += 1
        head.append(h if seen[h] == 1 else f"{h} ({seen[h]})")
    return [dict(zip(head, r)) for r in rows[hi + 1:] if any(v not in (None, "") for v in r)]


def n(v):
    try:
        f = float(v) if v not in (None, "") else None
        return f if f else None  # formula cells read 0 when the input is blank
    except (TypeError, ValueError):
        return None


def r1(v, k=1):
    return round(v, k) if v is not None else None


def s(v, lim=None):
    t = "" if v is None else str(v).strip()
    return t[:lim] if lim else t


def slug(name):
    t = unicodedata.normalize("NFKD", re.sub(r"\(.*?\)", "", name.lower())).encode("ascii", "ignore").decode()
    t = re.sub(r"[^\w\s-]", "", t).strip()
    return re.sub(r"[\s_-]+", "-", t)[:60] or "co"


SUFFIX = r"\b(the|co|cos|company|companies|corporation|corp|inc|llc|ltd|group|ag|plc|spa|s p a|nv|sa|holdings?|international)\b"


def key(name):
    """Comparable core of a company name: 'The Quikrete Companies' and 'Quikrete Cos.' both give 'quikrete'."""
    t = unicodedata.normalize("NFKD", s(name).lower()).encode("ascii", "ignore").decode()
    t = re.split(r"\s*\(|\s+/\s+", t)[0]
    t = re.sub(r"\s*(&|and)\s+sons?\b", " ", t.replace(".", " "))
    t = re.sub(SUFFIX, " ", t)
    return re.sub(r"[^a-z0-9]", "", t)


ALIASES = {"itw": "illinoistoolworks"}  # prior name -> Master key

# Product family of a category (first match wins).
CAT_SECTOR = [
    (r"hand (cleaner|wipe|soap|care)|skin|barrier cream|sunscreen", "hand-care"),
    (r"disinfect|sanitiz|insecticide|wasp|hornet|pest|rodent|odor|deodor", "sanitation"),
    (r"coil clean|ice machine|refrigera|vacuum pump|leak detect|hvac|condensate|duct seal|duct mastic|nitrogen", "hvac"),
    (r"pvc|cpvc|abs cement|solvent cement|pipe (thread|joint|dope)|thread seal|drain|septic|plumb|flux.*(solder|plumb)|boiler|hydronic|descal", "plumbing"),
    (r"wire pull|dielectric|anti.?oxid|joint compound|electrical|contact clean|fiber optic|arc|insulating|splic", "electrical"),
    (r"weld|spatter|cutting (fluid|oil)|tapping|layout|cold galv|soldering|brazing|metalworking|coolant", "welding"),
    (r"brake|carb|throttle|motor oil|oil additive|fuel|radiator|cooling system|antifreeze|body filler|automotive|fleet|def\b|diesel|transmission", "automotive"),
    (r"grout|mortar|concrete|cure|curing|form release|waterproof|firestop|anchor|masonry|bonding agent|tile|stucco|patch", "construction"),
    (r"water treat|dechlor|wastewater|test kit|reagent|chlorin|lift station", "water-ww"),
    (r"paint|marker|crayon|coating|primer|galvaniz|stain|lacquer", "paints"),
    (r"lubric|(?<!de)grease|penetr|oil|anti.?seize|silicone spray|rust (prevent|inhib)|corrosion", "lubricants"),
    (r"adhesive|sealant|caulk|thread ?lock|epoxy|glue|gasket|retaining|cement|mastic|foam", "adhesives"),
    (r"clean|degreas|solvent|remover|towel|wipe|glass|janitor", "cleaners"),
]


def cat_sector(cat):
    for pat, sid in CAT_SECTOR:
        if re.search(pat, cat or "", re.I):
            return sid
    return "other"


def split_prior(raw):
    """'CRC Industries (Berwind)' -> ('CRC Industries', 'Berwind'); 'Valvoline / Saudi Aramco' -> ('Valvoline', 'Saudi Aramco')."""
    raw = s(raw)
    m = re.match(r"^(.*?)\s*\((.+)\)\s*$", raw)
    if m:
        return m.group(1).strip(), m.group(2).strip()
    if " / " in raw:
        a, b = raw.split(" / ", 1)
        return a.strip(), b.strip()
    return raw, ""


brands = [b for b in table("Brands") if s(b.get("Brand"))]
products = table("Products")
dists = {s(d["Distributor"]): d for d in table("Distributors") if s(d.get("Distributor"))}
master = [m for m in table("Master") if s(m.get("Owner (ultimate)"))]
cos = [c for c in table("Companies") if s(c.get("Company"))]
# v24+: the Targets tab carries each owner's researched channel split (retail share, key retail customers, pro revenue)
targets_by_key = {}
if "Targets" in wb.sheetnames:
    for r in table("Targets"):
        if s(r.get("Owner")):
            targets_by_key.setdefault(key(s(r["Owner"])), r)


def n0(v):
    """Like n(), but a real 0 stays 0 (a 0% retail share is information)."""
    try:
        return float(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


# v28+: the Margin Model tab. A = inputs (anchor margin, gross margin, capture, haircuts, clips); B = EBITDA margin
# evidence from public comps; C = implied EBITDA margin per owner (anchor + scale haircut + price uplift).
margin_doc, mm_by_key = None, {}
if "Margin Model" in wb.sheetnames:
    sec, hdr, A, B = None, None, {}, []
    for row in wb["Margin Model"].iter_rows(values_only=True):
        c0 = str(row[0]).strip() if row[0] is not None else ""
        if c0[:3] in ("A. ", "B. ", "C. ", "D. "):
            sec, hdr = c0[0], None
            continue
        if not c0:
            continue
        if sec == "A":
            if isinstance(row[1], (int, float)):
                A[c0] = row[1]
        elif sec in ("B", "C") and hdr is None:
            hdr = [str(x).strip() if x is not None else "" for x in row]
        elif sec in ("B", "C"):
            r = dict(zip(hdr, row))
            if sec == "B":
                B.append({k: v for k, v in {"co": s(r.get("Company / source"), 80), "seg": s(r.get("Segment"), 120), "yr": r.get("Year"), "use": s(r.get("Use (Y/N/Ref)")),
                          "m": r1(n0(r.get("EBITDA margin")), 3), "ebit": r1(n0(r.get("EBIT margin")), 3), "gm": r1(n0(r.get("Gross margin")), 3),
                          "rev": r1(n(r.get("Revenue ($M)"))), "type": s(r.get("Type"), 40), "cat": s(r.get("Category"), 60), "note": s(r.get("Note"), 320),
                          "conf": s(r.get("Confidence")), "url": s(r.get("Source URL"), 300)}.items() if v not in ("", None)})
            else:
                mm_by_key[key(c0)] = {k: v for k, v in {"hc": r1(n0(r.get("Scale haircut")), 3), "up": r1(n0(r.get("Price uplift")), 4), "p": r1(n0(r.get("Price index (P)")), 3),
                    "pu": r1(n0(r.get("Index used (clipped)")), 3), "n": r1(n0(r.get("# priced items")), 0), "nb": r1(n0(r.get("# priced brands")), 0),
                    "conf": s(r.get("Best index confidence")), "lfl": s(r.get("Like-for-like pair data?")), "cat": s(r.get("Main priced category")),
                    "m": r1(n0(r.get("Implied EBITDA margin")), 4), "e": r1(n0(r.get("Implied EBITDA ($M)")))}.items() if v not in ("", None)}
    g = lambda pre: next((v for k, v in A.items() if k.startswith(pre)), None)
    margin_doc = {"anchor": g("Anchor EBITDA margin used"), "gm": g("Gross margin used"), "capture": g("Capture of price premium"),
                  "comps_med": g("Median EBITDA margin - comps"), "gm_med": g("Median gross margin - comps"),
                  "hc": [g("Scale haircut: revenue >= $250M"), g("Scale haircut: $50-250M"), g("Scale haircut: $10-50M"), g("Scale haircut: <$10M")],
                  "pclip": [g("Price index floor"), g("Price index cap")], "mclip": [g("Implied margin floor"), g("Implied margin cap")], "comps": B}
    print("margin model:", len(B), "comps,", sum(1 for b in B if b.get("use") == "Y"), "in the anchor,", len(mm_by_key), "owners with an implied margin")
deal_rows = [d for d in table("Deals") if s(d.get("Target"))]
profiles = {s(p.get("ID")): p for p in table("Profiles") if s(p.get("Company"))}
sources = [x for x in table("Sources") if s(x.get("Company"))]

# The same ultimate owner sometimes appears twice in Master under different IDs (two divisions of
# DuPont, say). Fold each duplicate into the one with the most brands; its company row rolls in.
canon = {}
for m in sorted(master, key=lambda m: (-(n(m.get("# brands")) or 0), n(m.get("Owner ID")) or 0)):
    canon.setdefault(key(m["Owner (ultimate)"]), s(m["Owner ID"]))
dupes = {s(m["Owner ID"]): canon[key(m["Owner (ultimate)"])] for m in master if canon[key(m["Owner (ultimate)"])] != s(m["Owner ID"])}
master = [m for m in master if s(m["Owner ID"]) not in dupes]
for c in cos:
    oid = s(c.get("Owner ID (rolls into)"))
    if oid in dupes:
        c["Owner ID (rolls into)"] = dupes[oid]
        if s(c.get("Role")) == "Owner":
            c["Role"], c["In owner figure? Owner/Y/N/X"] = "Merged duplicate", "Y"

# A Master row named "Unknown" is a bucket of brands whose maker wasn't found, not a company: its brands stay unowned.
placeholder = {s(m["Owner ID"]) for m in master if key(m["Owner (ultimate)"]) in ("unknown", "unidentified", "notidentified")}
master = [m for m in master if s(m["Owner ID"]) not in placeholder]

co_by_id = {s(c["ID"]): c for c in cos}
owner_row = {s(c["Owner ID (rolls into)"]): c for c in cos if s(c.get("Role")) == "Owner"}

# SKUs per brand per distributor, from the raw scrape.
sku_by = collections.defaultdict(collections.Counter)
for p in products:
    if s(p.get("Brand")):
        sku_by[s(p["Brand"])][s(p.get("Distributor"))] += 1

# ---- owners: verified (Master), then unverified priors that match no Master owner
owners = collections.OrderedDict()  # owner key -> {name, master, verified, brands, prior_parent}
ids = {}
for m in master:
    oid = s(m["Owner ID"])
    owners[oid] = {"name": s(m["Owner (ultimate)"]), "master": m, "verified": True, "brands": [], "parent": ""}
name_key = {}  # owner names only, for matching deal acquirers
for oid, o in owners.items():
    name_key.setdefault(key(o["name"]), oid)
by_key = dict(name_key)  # plus parent names, for merging priors into Master owners
for oid, o in owners.items():
    hero_parent = key(s(o["master"].get("Ultimate parent / sponsor")))
    if hero_parent and len(hero_parent) >= 5:
        by_key.setdefault(hero_parent, oid)
    for part in re.findall(r"\(([^)]+)\)", o["name"]):  # "Victor Technologies (ESAB)" is also ESAB
        if len(key(part)) >= 3:
            by_key.setdefault(key(part), oid)


def prior_owner(k):
    """A Master owner for an unverified prior name: exact key, else the one owner whose name starts with it."""
    if k in by_key:
        return by_key[k]
    hits = {o for nk, o in name_key.items() if len(k) >= 3 and nk.startswith(k)}
    return hits.pop() if len(hits) == 1 else None

unlinked = []
for b in brands:
    cid = s(b.get("Company ID"))
    if cid and cid in co_by_id:
        oid = s(co_by_id[cid]["Owner ID (rolls into)"])
        (owners[oid]["brands"] if oid in owners else unlinked).append(b)
        continue
    raw = s(b.get("Likely owner (PRIOR - unverified)"))
    if not raw:
        unlinked.append(b)
        continue
    name, parent = split_prior(raw)
    k = ALIASES.get(key(name), key(name))
    oid = prior_owner(k)
    if not oid:
        oid = "prior:" + k
        by_key[k] = oid
        owners[oid] = {"name": name, "master": {}, "verified": False, "brands": [], "parent": parent,
                       "otype": s(b.get("Owner type (PRIOR - unverified)"))}
    owners[oid]["brands"].append(b)
    b["_prior"] = True

for oid, o in owners.items():
    cid = slug(o["name"])
    while cid in ids.values():
        cid += "-2"
    ids[oid] = cid

# ---- deals: an owner's own acquisitions go on its M&A tab; ownership changes and the rest go to meta/deals
acq = collections.defaultdict(list)
other_deals, seen_deals = [], set()
pe_entry = {}
for d in sorted(deal_rows, key=lambda d: s(d.get("Date")) or s(d.get("Year")), reverse=True):
    tgt = co_by_id.get(s(d.get("Target ID")))
    date = s(d.get("Date")) or s(d.get("Year"))
    dk = (key(d["Target"])[:12], date[:4])
    if dk in seen_deals:  # the ledger repeats a few deals with slightly different names
        continue
    seen_deals.add(dk)
    row = {"date": date, "target": s(d["Target"]), "deal_type": s(d.get("Deal type")), "what_it_does": s(tgt.get("What it mostly is (blunt)"), 300) if tgt else "",
           "what_it_added": s(d.get("Target brands")), "value_usd_m": n(d.get("Value ($M, if disclosed)")), "source_url": s(d.get("Source")) if s(d.get("Source")).startswith("http") else "",
           "source": s(d.get("Source")), "confidence": "High" if s(d.get("Source")).startswith("http") else "Medium"}
    owner = s(tgt["Owner ID (rolls into)"]) if tgt else ""
    if tgt and s(tgt.get("Role")) != "Owner" and owner in owners:
        acq[owner].append(row)  # the owner bought one of the companies now rolled into it
        continue
    if tgt and s(tgt.get("Role")) == "Owner" and name_key.get(key(d.get("Acquirer"))) == owner:
        acq[owner].append(row)  # the owner bought the business that now carries its name (Clorox buying GOJO)
        continue
    if not tgt:  # match the acquirer's name to an owner
        a = name_key.get(key(d.get("Acquirer")))
        if a and a in owners:
            acq[a].append(row)
            continue
    # ownership change of an owner, or a deal between companies outside the universe
    other = {**row, "acquirer": s(d.get("Acquirer")), "target_owner_id": ids.get(owner) if owner in owners else None}
    other_deals.append(other)

# A PE-backed owner's entry: the latest deal on it (or a business now inside it) where the acquirer names its sponsor.
for d in sorted(deal_rows, key=lambda d: s(d.get("Date")) or s(d.get("Year")), reverse=True):
    tgt = co_by_id.get(s(d.get("Target ID")))
    owner = s(tgt["Owner ID (rolls into)"]) if tgt else next((o for k, o in name_key.items() if len(k) >= 6 and key(d["Target"])[:8] == k[:8]), "")
    m = owners[owner]["master"] if owner in owners else {}
    sp, acq_by = key(m.get("PE sponsor")), key(d.get("Acquirer"))  # key() drops anything in brackets
    if owner and owner not in pe_entry and s(m.get("PE-backed")) == "Y" and len(sp) >= 3 and len(acq_by) >= 3 and (acq_by.startswith(sp[:6]) or sp.startswith(acq_by[:6])):
        pe_entry[owner] = {"date": s(d.get("Date")) or s(d.get("Year")), "value_usd_m": n(d.get("Value ($M, if disclosed)"))}

ev_by_co = collections.defaultdict(list)
for x in sources:
    ev_by_co[s(x["ID"])].append({"metric": s(x.get("Metric")), "value": s(x.get("Value"), 200), "year": s(x.get("Year")),
                                 "url": s(x.get("URL")), "quote": s(x.get("Quote"), 300)})

# ---- channels: which distributors are pro channels, and which end market each serves
RETAIL = {"Tractor Supply", "Do it Best / Ace / True Value (coops)", "Rockler / Woodcraft", "Pool Geek", "INYOPools", "TCP Global",
          "The Home Depot (Pro desk)", "West Marine (Pro)", "AutoZone (Commercial)", "O'Reilly Auto Parts (Pro)",
          "Advance Auto Parts (Carquest)", "Leslie's Pro", "Amazon Business"}
BROADLINE = {"W.W. Grainger", "McMaster-Carr", "Fastenal", "MSC Industrial Supply", "Applied Industrial Technologies", "Motion Industries",
             "Global Industrial", "Zoro", "Kimball Midwest", "Lawson Products (DSG)", "ORS Nasco", "Uline", "Home Depot Pro (incl. HD Supply)"}
MARKET = {"Industrial MRO": "Industrial MRO", "Industrial MRO (wholesale)": "Industrial MRO", "Safety & PPE": "Industrial MRO",
          "Oil & Gas / Energy": "Industrial MRO", "Plumbing": "Plumbing", "HVAC/R": "HVAC/R", "Electrical": "Electrical",
          "Facilities Maintenance": "Facilities & jan-san", "Janitorial & Sanitation": "Facilities & jan-san", "Facilities / Shipping": "Facilities & jan-san",
          "Construction / Contractor": "Construction", "Roofing & Exteriors": "Construction", "Drywall & Interiors": "Construction",
          "Flooring & Tile": "Construction", "Paint & Coatings": "Construction", "Woodworking & Cabinetry": "Construction",
          "Automotive Aftermarket": "Auto & fleet", "Automotive Refinish": "Auto & fleet", "Heavy Duty / Fleet": "Auto & fleet",
          "Welding & Fabrication": "Welding", "Water & Wastewater": "Water & pool", "Pool & Irrigation": "Water & pool", "Pool & Spa": "Water & pool",
          "Marine": "Marine & aviation", "Aviation MRO": "Marine & aviation", "Agriculture": "Farm & landscape", "Landscape & Turf": "Farm & landscape",
          "Electronics Assembly": "Electronics & lab", "Lab & Cleanroom": "Electronics & lab", "Hardware": "Hardware", "Hardware (wholesale)": "Hardware"}


def channel(d):
    return "retail" if d in RETAIL else "broadline" if d in BROADLINE else "specialty"


def market(d):
    return MARKET.get(s(dists.get(d, {}).get("Primary end market")), "Other")


# v20+: the workbook's own channel type per distributor (Pro, Mixed pro-DIY, Retail store, DIY / enthusiast online)
CTYPE = {"Pro": "pro", "Mixed pro-DIY": "mixed", "Retail store": "retail", "DIY / enthusiast online": "diy"}


def ctype(d):
    return CTYPE.get(s(dists.get(d, {}).get("Channel type")), "")


FAMILY = {"lubricants": "Lubricants & penetrants", "adhesives": "Adhesives & sealants", "cleaners": "Cleaners & degreasers",
          "hand-care": "Hand care", "sanitation": "Disinfectants & pest", "paints": "Paints & markers", "plumbing": "Plumbing chemicals",
          "hvac": "HVAC/R chemicals", "electrical": "Electrical chemicals", "welding": "Welding & metalworking", "automotive": "Auto & fleet chemicals",
          "construction": "Construction chemicals", "water-ww": "Water treatment", "other": "Other"}

# ---- what a worker can hold: format, then pack size for bottles, cans and powders
HANDHELD = {"Aerosol", "Squeeze tube", "Cartridge", "Wipe", "Marker/Paint stick", "Kit"}
SIZED = {"Bottle/Liquid", "Paste/Can", "Powder"}
BULK_WORDS = re.compile(r"\b(drum|pail|tote|bulk|55[- ]?gal|5[- ]?gal|2\.5[- ]?gal|bag-in-box)\b", re.I)


def pack_ok(size, name):
    """True when a bottle, can or powder is hand-sized: up to 1 gal / 4 L, or 10 lb / 5 kg."""
    t = f"{size} {name}".lower()
    if BULK_WORDS.search(t):
        return False
    m = re.search(r"(\d+(?:\.\d+)?)\s*-?\s*(fl\.?\s*oz|oz|gal(?:lon)?s?|qt|quarts?|pt|pints?|lbs?|pounds?|ml|l\b|liters?|litres?|kg|g\b)", t)
    if not m:
        return True  # no size listed: most scraped bottles and cans are hand-sized
    v, u = float(m.group(1)), m.group(2)
    if u.startswith("gal"):
        return v <= 1
    if u.startswith(("lb", "pound")):
        return v <= 10
    if u == "kg":
        return v <= 5
    if u in ("l",) or u.startswith(("liter", "litre")):
        return v <= 4
    if u.startswith("qt") or u.startswith("quart"):
        return v <= 4
    return True  # oz, pt, ml, g


DURABLE = re.compile(r"dispens|\btool|\bgun\b|pump\b|sprayer|brush|equipment|machine|gauge|fitting|valve|hose|bucket|mop\b|applicator|"
                     r"trowel|tape measure|meter\b|cabinet|can opener|plug\b|torch|regulator|nozzle|wrench|blade|battery|charger|light|"
                     r"filter housing|stand\b|cart\b|holder", re.I)


def handheld(p):
    f = s(p.get("Form factor"))
    if f in HANDHELD:
        return True
    if f in SIZED:
        return pack_ok(s(p.get("Size / pack")), s(p.get("Product name")))
    return False


def consumable(p):
    if s(p.get("Form factor")) in HANDHELD | SIZED - {"Kit"}:
        return not DURABLE.search(s(p.get("Product category")))
    return not DURABLE.search(f"{s(p.get('Product category'))} {s(p.get('Product name'))}") and s(p.get("Form factor")) != "Other"


prod_by_brand = collections.defaultdict(list)
for p in products:
    if s(p.get("Brand")):
        prod_by_brand[s(p["Brand"])].append(p)

# ---- categories (v12+): the workbook gives each brand one normalized category in a segment. A brand's SKUs often
# span several, so each SKU's raw category is mapped to the normalized category most brands using that raw label
# sit in; a brand then counts in every category it has product in, not only its primary one.
comp_rows = [r for r in table("Competition") if s(r.get("Category (normalized)")) and not s(r["Category (normalized)"]).startswith("Total")] \
    if "Competition" in wb.sheetnames else []
SEG_OF = {s(r["Category (normalized)"]): s(r.get("Segment")) or "Unclassified" for r in comp_rows}
brand_ncat = {s(b["Brand"]): s(b.get("Category (normalized)")) for b in brands}
rawmap = collections.defaultdict(collections.Counter)
for p in products:
    nc = brand_ncat.get(s(p.get("Brand")))
    if nc and nc != "Unclassified" and s(p.get("Product category")):
        rawmap[s(p["Product category"]).lower()][nc] += 1


def sku_ncat(p, fallback):
    c = rawmap.get(s(p.get("Product category")).lower())
    if c:
        (top, k), tot = c.most_common(1)[0], sum(c.values())
        if tot >= 2 and k / tot >= 0.5:
            return top
    return fallback or "Unclassified"

REV_MIX = {"trade": ["% Industrial MRO", "% Plumbing", "% HVAC/R", "% Electrical", "% Construction & Contractor", "% Facilities & Janitorial",
                     "% Automotive & Fleet", "% Welding & Fabrication", "% Water & Wastewater"],
           "fmt": ["% Aerosol", "% Squeeze tube", "% Cartridge", "% Bottle/Liquid", "% Wipe", "% Marker/Paint stick", "% Paste/Can"]}
OWNER_CLASS = {"PE-backed": "pe", "Founder / family": "family", "Private (family)": "family", "ESOP / employee": "esop",
               "Subsidiary of private co": "sub_private", "Subsidiary of public co": "public", "Public": "public",
               "Unknown": "private", "Private": "private", "PE": "pe"}

out_co, out_dt, judge_in = {}, {}, []
for oid, o in owners.items():
    cid, m = ids[oid], o["master"]
    bs = sorted(o["brands"], key=lambda b: -(n(b.get("# SKUs scraped")) or 0))
    c = owner_row.get(oid, {}) if o["verified"] else {}
    rolled = [x for x in cos if s(x.get("Owner ID (rolls into)")) == oid and s(x.get("Role")) != "Owner"]
    raw_type = s(m.get("Owner type")) or o.get("otype", "")
    oclass = OWNER_CLASS.get(raw_type, "private")
    if oclass == "private" and s(m.get("PE-backed")) == "Y":
        oclass = "pe"

    fmt, mkt, dst, fam, brd, seg = (collections.Counter() for _ in range(6))
    hand = cons = pro = 0
    brand_rows, max_bd = [], 0
    for b in bs:
        ps = prod_by_brand.get(s(b["Brand"]), [])
        bd = collections.Counter(s(p.get("Distributor")) for p in ps if s(p.get("Distributor")))
        max_bd = max(max_bd, sum(1 for d in bd if channel(d) != "retail"))
        for p in ps:
            d = s(p.get("Distributor"))
            fmt[s(p.get("Form factor")) or "Other"] += 1
            mkt[market(d)] += 1
            dst[d] += 1
            nc = sku_ncat(p, s(b.get("Category (normalized)")))
            fam[nc] += 1
            seg[SEG_OF.get(nc, "Unclassified")] += 1
            brd[s(b["Brand"])] += 1
            hand += handheld(p)
            cons += consumable(p)
            pro += channel(d) != "retail"
        brand_rows.append({"name": s(b["Brand"]), "cat": s(b.get("Primary product category")), "hero": s(b.get("Known for (hero product)"), 140),
                           "form": s(b.get("Primary form factor")), "skus": len(ps), "ncat": s(b.get("Category (normalized)")),
                           "dists": [d for d, _ in bd.most_common()], "mkts": sorted({market(d) for d in bd}),
                           "verified": not b.get("_prior")})
    skus = sum(fmt.values())
    nd_pro = sum(1 for d in dst if channel(d) != "retail")
    rev = n(c.get("Total company rev ($M)")) or n(m.get("TOTAL PLATFORM REV ($M)"))
    pe = pe_entry.get(oid, {})
    ptxt = s((profiles.get(s(c.get("ID"))) or {}).get("Profile"))
    paras = [p.strip() for p in re.split(r"\n\s*\n", ptxt) if p.strip()]
    body = [p for p in paras if not p.startswith(("# ", "**", "Brands:"))]
    blunt = s(m.get("What it mostly is (blunt)"))
    what = (body[0] if body else blunt)[:600]
    summary = {
        "name": o["name"], "verified": o["verified"], "oclass": oclass, "otype": raw_type or "Unknown",
        "sponsor": s(m.get("PE sponsor")) or (o["parent"] if oclass == "pe" else ""),
        "parent": s(m.get("Ultimate parent / sponsor")) if o["verified"] else o["parent"],
        "pe_year": pe.get("date", "")[:4], "pe_ev": pe.get("value_usd_m"), "rollup": s(m.get("PE roll-up")) == "Y",
        "model": s(m.get("Business model")), "trade": s(m.get("Primary trade")) or (s(bs[0].get("Primary trade")) if bs else ""),
        "category": s(m.get("Primary product category")) or (s(bs[0].get("Primary product category")) if bs else ""),
        "blunt": blunt, "hq": ", ".join(x for x in [s(c.get("HQ city")), s(c.get("HQ state")) or s(m.get("HQ state"))] if x),
        "web": s(c.get("Website")), "rev": r1(rev), "rev_year": s(c.get("Revenue year")), "rev_basis": s(c.get("Revenue basis / evidence"), 200),
        **(lambda tg: {k: v for k, v in {"ret_pct": r1(n0(tg.get("Retail share - company-level (%)")), 1),
            "ch_class": s(tg.get("Channel classification (best available)")), "ch_conf": s(tg.get("Split confidence")),
            "prorev": r1(n(tg.get("Pro-channel revenue est. ($M)"))),
            "im_m": r1(n0(tg.get("Implied EBITDA margin (Margin Model)")), 4), "im_e": r1(n(tg.get("Implied EBITDA ($M)"))),
            "lfl": r1(n0(tg.get("Price index vs peers (like-for-like)")), 3), "lfl_conf": s(tg.get("Price index confidence"))}.items() if v not in ("", None)})(targets_by_key.get(key(o["name"]), {})),
        **({"mm": mm_by_key[key(o["name"])]} if key(o["name"]) in mm_by_key else {}),
        "emp": r1(n(c.get("Employees (total)")), 0), "emp_band": s(c.get("Employees band")), "size_basis": s(c.get("Size basis")),
        "size_conf": s(c.get("Size confidence")), "linkedin": s(c.get("LinkedIn URL")), "li_emp": r1(n(c.get("LinkedIn employees (baseline Oct-2026)")), 0),
        "pro_rev": r1(n(c.get("Pro brand rev ($M, US)")) or n(m.get("PRO BRAND REV ($M, pro forma)"))),
        "ebitda": r1(n(c.get("EBITDA reported ($M)")) or n(m.get("PLATFORM EBITDA ($M)"))), "ebitda_rep": True if n(c.get("EBITDA reported ($M)")) else None, "margin": r1(n(m.get("EBITDA margin")), 3),
        "rev_mix": {g: {k[2:]: r1(n(m.get(k)) * (100 if n(m.get(k)) <= 1 else 1)) for k in ks if n(m.get(k))} for g, ks in REV_MIX.items()
                    if any(n(m.get(k)) for k in ks)},
        "conf": s(m.get("Confidence")) or "Low",
        "brands": [x["name"] for x in brand_rows] or [b.strip() for b in s(m.get("Hero brands")).split(",") if b.strip()],
        "skus": skus, "acq": len(acq.get(oid, [])), "rolled": [s(x["Company"]) for x in rolled],
        "mix": {"fmt": dict(fmt), "mkt": dict(mkt), "dist": dict(dst.most_common(15)), "cat": dict(fam), "seg": dict(seg), "brand": dict(brd.most_common(15)),
                "ch": dict(collections.Counter(ctype(d) or "unknown" for d, v in dst.items() for _ in range(v)))},
        "m": {"hand": r1(hand / skus, 3) if skus else None, "cons": r1(cons / skus, 3) if skus else None, "pro": r1(pro / skus, 3) if skus else None,
              "nd": len(dst), "ndp": nd_pro, "maxbd": max_bd,
              "broad": r1(sum(v for d, v in dst.items() if channel(d) == "broadline") / skus, 3) if skus else None},
    }
    summary = {k: v for k, v in summary.items() if v not in ("", None, [], {})}
    out_co[cid] = summary
    out_dt[cid] = {k: v for k, v in {
        "what": what, "size_note": s(c.get("Size sources / note"), 600),
        "ch_basis": s(targets_by_key.get(key(o["name"]), {}).get("Channel-split basis / key retail customers"), 600), "profile": "\n\n".join(paras[1:]) if paras and paras[0].startswith("# ") else ptxt,
        "brands": brand_rows, "acq": acq.get(oid, []),
        "evidence": [e for x in [c] + rolled for e in ev_by_co.get(s(x.get("ID")), [])][:30],
    }.items() if v not in ("", None, [])}
    judge_in.append({"id": cid, "name": o["name"], "owner": raw_type, "parent": summary.get("parent", ""), "model": summary.get("model", ""),
                     "category": summary.get("category", ""), "what": what[:400], "brands": [f"{x['name']} ({x['cat']}; {x['form']}; {x['skus']} SKUs)" for x in brand_rows[:12]],
                     "revenue_usd_m": summary.get("rev"), "skus": skus, "handheld_share": summary["m"]["hand"], "pro_channel_share": summary["m"]["pro"],
                     "pro_distributors": nd_pro})

# ---- every deal, with owners resolved where the ledger allows
deals_out, seen = [], set()
for d in sorted(deal_rows, key=lambda d: s(d.get("Date")) or s(d.get("Year")), reverse=True):
    date = s(d.get("Date")) or s(d.get("Year"))
    dk = (key(d["Target"])[:12], date[:4])
    if dk in seen:
        continue
    seen.add(dk)
    tgt = co_by_id.get(s(d.get("Target ID")))
    t_owner = s(tgt["Owner ID (rolls into)"]) if tgt else ""
    a_owner = name_key.get(key(d.get("Acquirer")))
    if not a_owner and tgt and s(tgt.get("Role")) != "Owner" and t_owner in owners:
        a_owner = t_owner  # the target now rolls into an owner, so that owner (or its predecessor) bought it
    self_buy = tgt is not None and s(tgt.get("Role")) == "Owner" and a_owner == t_owner
    kind = "addon" if a_owner in owners and (not tgt or s(tgt.get("Role")) != "Owner" or self_buy) else "ownership" if t_owner in owners else "other"
    src = s(d.get("Source"))
    deals_out.append({k: v for k, v in {
        "date": date, "year": int(date[:4]) if date[:4].isdigit() else None, "acquirer": s(d.get("Acquirer")),
        "acq_id": ids.get(a_owner) if a_owner in owners else None, "target": s(d["Target"]), "tgt_id": ids.get(t_owner) if t_owner in owners else None,
        "brands": s(d.get("Target brands")), "value": n(d.get("Value ($M, if disclosed)")), "type": s(d.get("Deal type")),
        "url": src if src.startswith("http") else "", "kind": kind}.items() if v not in ("", None)})

# ---- every brand: its owner (verified, unverified prior, Claude's guess, or unknown), products and channels
BOF = os.path.join(os.path.dirname(os.path.abspath(__file__)), "brand_owners.json")
guesses = json.load(open(BOF)) if os.path.exists(BOF) else {}
GUESS_CLASS = {"PE": "pe", "Family": "family", "Public": "public", "Subsidiary of public": "public",
               "Subsidiary of private": "sub_private", "Private (unknown)": "private", "Unknown": "private",
               "Founder / family": "family", "PE-backed": "pe", "Subsidiary of public co": "public", "Subsidiary of private co": "sub_private",
               "Public": "public", "ESOP": "esop"}
brand_owner = {s(b["Brand ID"]): (oid, "prior" if b.get("_prior") else "verified") for oid, o in owners.items() for b in o["brands"]}
maker = collections.defaultdict(collections.Counter)
for p in products:
    if s(p.get("Manufacturer as listed")):
        maker[s(p.get("Brand"))][re.sub(r"[®™]", "", s(p["Manufacturer as listed"]))] += 1
out_br, out_sku, wb_brand_id = {}, {}, {}


def parse_competitors(txt):
    """'Weld-On (IPS) (PE: Centerbridge Partners); Spears (Family)' -> [["Weld-On (IPS)", "PE: Centerbridge Partners"], ["Spears", "Family"]]"""
    out = []
    for x in (txt or "").split(";"):
        x = x.strip()
        if not x:
            continue
        if x.endswith(")") and " (" in x:
            out.append([x[:x.rfind(" (")].strip(), x[x.rfind(" (") + 2:-1].strip()])
        else:
            out.append([x, ""])
    return out

# v9 added a triage pass on every brand (bucket, owner guess, owner type, confidence)
TRIAGE = {"ACTIONABLE_CANDIDATE": "actionable", "UNKNOWN_BRAND": "unknown", "SUBSIDIARY_KNOWN": "sub_known",
          "SUBSIDIARY_PUBLIC": "sub_public", "OUT_OF_SCOPE_OTHER": "out", "NOT_CHEMICAL": "not_chem",
          "FOREIGN_MAKER": "foreign", "HOUSE_BRAND": "house", "CONSUMER_ONLY": "consumer", "SUBSIDIARY_OTHER_LARGE": "sub_large"}
guess_by_name = {key(k) or k.lower(): v for k, v in guesses.items()}
for b in brands:
    name = s(b["Brand"])  # brand ids come from the name: the workbook renumbers its Brand IDs between versions
    bid = slug(name) or "brand"
    while bid in out_br:
        bid += "-2"
    ps = prod_by_brand.get(name, [])
    bd = collections.Counter(s(p.get("Distributor")) for p in ps if s(p.get("Distributor")))
    fm = collections.Counter(s(p.get("Form factor")) or "Other" for p in ps)
    ncat = s(b.get("Category (normalized)")) or "Unclassified"
    cats = collections.Counter(sku_ncat(p, ncat) for p in ps)
    rec = {"name": name, "cat": s(b.get("Primary product category")), "ncat": ncat, "seg": s(b.get("Segment")) or SEG_OF.get(ncat, "Unclassified"),
           "cats": dict(cats.most_common()) if len(cats) > 1 else {},
           "wc": parse_competitors(s(b.get("Top competitors - same category, ranked by distributor reach (owner type; ? = unverified)"))),
           "wcmix": s(b.get("Competitor owner mix (top 10: PE / family+ESOP / public+public sub / other)")),
           "price": {k: v for k, v in {"pos": s(b.get("Price position")), "idx": r1(n(b.get("Price index vs category")), 2),
                     "ppo": r1(n(b.get("Avg $/oz - brand (main category)")), 3), "cpo": r1(n(b.get("Avg $/oz - category")), 3),
                     "n": r1(n(b.get("# priced items (main category)")), 0)}.items() if v not in ("", None)},
           "review": {k: v for k, v in {"site": s(b.get("Review site")), "prod": s(b.get("Most-reviewed product"), 140),
                      "n": r1(n(b.get("# reviews (that product)")), 0), "rating": r1(n(b.get("Avg rating (that product)")), 2),
                      "url": s(b.get("Review URL"), 300), "note": s(b.get("Review note"), 200)}.items() if v not in ("", None)},
           "wb_owner": s(b.get("Owner / sponsor (best available)")), "wb_otype": s(b.get("Owner type (best available)")),
           "wb_basis": s(b.get("Owner basis")),
           "hero": s(b.get("Known for (hero product)"), 160), "form": s(b.get("Primary form factor")), "fmt": dict(fm), "skus": len(ps),
           "dists": [d for d, _ in bd.most_common()], "mkts": sorted({market(d) for d in bd}), "scope": s(b.get("Scope flag")),
           "ch": dict(collections.Counter(ctype(s(p.get("Distributor"))) or "unknown" for p in ps)),
           "ret": r1(n0(b.get("Retail share used (%)")), 1), "ret_b": {"owner research": "o", "brand scrape": "s"}.get(s(b.get("Retail share basis")), ""),
           "rscreen": {k: v for k, v in {"tier": s(b.get("Retail screen tier (big-box/mass search)")), "hits": s(b.get("Retail screen hits"), 200),
                       "signals": s(b.get("Consumer signals (screen)"), 300)}.items() if v},
           "maker": [m for m, _ in maker[name].most_common(2)], "tri": TRIAGE.get(s(b.get("Triage bucket (single-source)"))),
           "tri_owner": s(b.get("Triage owner guess")), "tri_type": s(b.get("Triage owner type guess")), "tri_conf": s(b.get("Triage confidence")), "rev": r1(n(b.get("Est. brand rev ($M)"))),
           "hand": r1(sum(handheld(p) for p in ps) / len(ps), 2) if ps else None}
    oid, status = brand_owner.get(s(b["Brand ID"]), (None, "unknown"))
    if oid:
        rec.update(owner_id=ids[oid], owner=owners[oid]["name"], status=status, oclass=out_co[ids[oid]]["oclass"])
    else:
        g = guesses.get(name) or guess_by_name.get(key(name) or name.lower()) or {}
        if not s(g.get("owner")) and s(b.get("Triage owner guess")) and s(b.get("Triage owner guess")).lower() not in ("unknown", "n/a"):
            g = {"owner": s(b["Triage owner guess"]), "type": s(b.get("Triage owner type guess")), "conf": s(b.get("Triage confidence")) or "Low",
                 "note": "The workbook's triage guess.", "src": "triage"}
        if s(g.get("owner")):
            hit = name_key.get(key(g["owner"])) or name_key.get(key(g.get("parent")))
            rec.update(status="guess", owner=s(g["owner"]), parent=s(g.get("parent")), oclass=GUESS_CLASS.get(s(g.get("type")), "private"),
                       sponsor=s(g.get("sponsor")), conf=s(g.get("conf")) or "Low", note=s(g.get("note"), 120), src=s(g.get("src")),
                       **({"owner_id": ids[hit]} if hit in owners else {}))
        else:
            rec.update(status="unknown", oclass="unknown")
    out_br[bid] = {k: v for k, v in rec.items() if v not in ("", None, [], {})}
    wb_brand_id[bid] = s(b["Brand ID"])
    out_sku[bid] = [[s(p.get("Product name"), 140), s(p.get("Form factor")), s(p.get("Size / pack"), 40), s(p.get("Distributor")),
                     s(p.get("Product URL"), 300), s(p.get("Product category"), 60), s(p.get("List price ($)")), s(p.get("Mfr part #"), 40)] for p in ps]

# ---- workbook competitor lists: link names to brand pages
bid_by_name = {}
for bid, r in out_br.items():
    bid_by_name.setdefault(key(r["name"]) or r["name"].lower(), bid)
for r in out_br.values():
    for e in r.get("wc", []):
        hit = bid_by_name.get(key(e[0]) or e[0].lower()) or bid_by_name.get(key(re.sub(r"\s*\(.*\)$", "", e[0])) or "")
        if hit:
            e.append(hit)


# ---- where Claude's owner guess and the workbook's triage guess name different companies: a verification worklist
def same_owner(a, b):
    a, b = key(a or ""), key(b or "")
    return bool(a and b and (a in b or b in a or a[:6] == b[:6]))


for r in out_br.values():
    w = r.get("wb_owner", "")
    if r.get("status") == "guess" and r.get("src") != "triage" and r.get("wb_basis", "").startswith("Triage") and w \
            and not re.match(r"unknown|n/a|none", w, re.I) and not same_owner(r.get("owner"), w) and not same_owner(r.get("parent"), w):
        r["odis"] = 1
print("owner disagreements (Claude guess vs workbook triage):", sum(1 for r in out_br.values() if r.get("odis")))

# ---- prices captured from distributor sites, per brand (loaded with the brand page)
out_q = collections.defaultdict(list)
if "Prices" in wb.sheetnames:
    for q in table("Prices"):
        bid = bid_by_name.get(key(s(q.get("Brand (mapped)"))) or s(q.get("Brand (mapped)")).lower())
        if bid:
            out_q[bid].append([s(q.get("Product"), 140), s(q.get("Size"), 30), r1(n(q.get("Pack qty")), 0), r1(n(q.get("Price ($)")), 2),
                               s(q.get("Price unit"), 12), r1(n(q.get("$ per oz")), 3), s(q.get("Use (Y/N)")) == "Y",
                               s(q.get("Excluded because"), 40), s(q.get("Distributor")), s(q.get("Product URL"), 300) or s(q.get("Source page"), 300),
                               s(q.get("Category (normalized)"))])

# ---- price ladder: each usable price indexed against its category's median $/oz for the same pack-size band (a gallon
# is far cheaper per oz than a 12 oz can, so raw $/oz would mostly measure pack mix). A brand's premium index in a
# category = the median of its prices' indices; pro and retail indices use only prices from those channel types.
def size_band(oz):
    return "<8" if oz < 8 else "8-32" if oz <= 32 else "32-128" if oz <= 128 else ">128"


def med(xs):
    xs, h = sorted(xs), len(xs) // 2
    return None if not xs else xs[h] if len(xs) % 2 else (xs[h - 1] + xs[h]) / 2


price_rows = []
for bid, rows in out_q.items():
    for q in rows:
        if q[6] and q[5] and q[10]:
            price_rows.append((bid, q[10], q[5], q[8]))
oz_by = {}
if "Prices" in wb.sheetnames:  # the oz-equivalent size, for the band
    for q in table("Prices"):
        b2 = bid_by_name.get(key(s(q.get("Brand (mapped)"))) or s(q.get("Brand (mapped)")).lower())
        if b2 and s(q.get("Use (Y/N)")) == "Y" and n(q.get("$ per oz")) and n(q.get("Size (oz-equiv)")):
            oz_by.setdefault((b2, s(q.get("Category (normalized)")), round(n(q["$ per oz"]), 3), s(q.get("Distributor"))), []).append(n(q["Size (oz-equiv)"]))
cells, cat_all = collections.defaultdict(list), collections.defaultdict(list)
priced = []
for bid, cat, ppo, dist in price_rows:
    ozs = oz_by.get((bid, cat, ppo, dist))
    if not ozs:
        continue
    band = size_band(ozs[0])
    cells[(cat, band)].append(ppo); cat_all[cat].append(ppo)
    priced.append((bid, cat, ppo, dist, band))
cell_med = {k: med(v) for k, v in cells.items() if len(v) >= 3}
cat_med = {k: med(v) for k, v in cat_all.items()}
pb = collections.defaultdict(lambda: {"i": [], "p": [], "pro": [], "ret": []})
for bid, cat, ppo, dist, band in priced:
    base = cell_med.get((cat, band)) or cat_med.get(cat)
    if not base:
        continue
    x = pb[(cat, bid)]; ix = ppo / base
    x["i"].append(ix); x["p"].append(ppo)
    (x["ret"] if ctype(dist) in ("retail", "diy") else x["pro"]).append(ix)
pricing_out = collections.defaultdict(dict)
for (cat, bid), x in pb.items():
    pricing_out[cat][bid] = [len(x["i"]), r1(med(x["i"]), 3), r1(med(x["p"]), 3), r1(min(x["i"]), 3), r1(max(x["i"]), 3),
                             r1(med(x["pro"]), 3) if x["pro"] else None, r1(med(x["ret"]), 3) if x["ret"] else None,
                             [sum(i <= 0.87 for i in x["i"]), sum(0.87 < i < 1.15 for i in x["i"]), sum(i >= 1.15 for i in x["i"])]]
# v28+: the market map's like-for-like index (Margin Model section D) replaces the pack-size index as the headline:
# log $/oz = category + distributor + category x form + category-specific size slope + a ridge-shrunk brand effect,
# so 1.00 = the median brand in the category on the same shelf, format and size. Our own per-price indices are kept
# for the spread and channel split, rescaled so their median equals the like-for-like index. Tiers follow the
# workbook: Good <= 0.80, Better, Best >= 1.25. Rows the model doesn't cover keep the pack-size index (no confidence).
LO, HI = 0.80, 1.25
lfl = {}
if "Margin Model" in wb.sheetnames:
    sec, hdr = None, None
    for row in wb["Margin Model"].iter_rows(values_only=True):
        c0 = str(row[0]).strip() if row[0] is not None else ""
        if c0[:3] in ("A. ", "B. ", "C. ", "D. "):
            sec, hdr = c0[0], None
            continue
        if sec != "D" or not c0:
            continue
        if hdr is None:
            hdr = [str(x).strip() if x is not None else "" for x in row]
            continue
        r = dict(zip(hdr, row))
        b2 = bid_by_name.get(key(c0)) or bid_by_name.get(c0.lower())
        if b2 and n0(r.get("Price index")):
            lfl[(s(r.get("Category")), b2)] = (n0(r["Price index"]), int(n0(r.get("# items")) or 0), int(n0(r.get("# distributors")) or 0),
                                               int(n0(r.get("# same-shelf peer brands")) or 0), (s(r.get("Confidence")) or "L")[:1])
tier3 = lambda xs: [sum(i <= LO for i in xs), sum(LO < i < HI for i in xs), sum(i >= HI for i in xs)]
n_lfl = 0
for (cat, bid), (ix, items, nd, npeer, conf) in lfl.items():
    x = pb.get((cat, bid))
    if x:
        k = ix / med(x["i"]); sc = lambda xs: [i * k for i in xs]
        pricing_out[cat][bid] = [len(x["i"]), r1(ix, 3), r1(med(x["p"]), 3), r1(min(x["i"]) * k, 3), r1(max(x["i"]) * k, 3),
                                 r1(med(x["pro"]) * k, 3) if x["pro"] else None, r1(med(x["ret"]) * k, 3) if x["ret"] else None, tier3(sc(x["i"])), conf, nd, npeer]
    else:
        pricing_out[cat][bid] = [items or 1, r1(ix, 3), None, r1(ix, 3), r1(ix, 3), None, None, tier3([ix] * (items or 1)), conf, nd, npeer]
    n_lfl += 1
for cat, bs in pricing_out.items():  # rows the model doesn't cover: pack-size index, re-tiered on the same cutoffs
    for bid, v in bs.items():
        if len(v) < 9:
            x = pb[(cat, bid)]; v[7] = tier3(x["i"]); v += [None, None, None]
print("like-for-like:", n_lfl, "brand-category rows from the Margin Model;", sum(1 for bs in pricing_out.values() for v in bs.values() if v[8] is None), "kept on the pack-size index")
pricing_meta = {cat: {"rows": len(cat_all.get(cat, [])), "med": r1(cat_med.get(cat), 3), "src": "lfl",
                      "bands": {b: r1(v, 3) for (c, b), v in cell_med.items() if c == cat}} for cat in pricing_out}
# what the ladder teaches, for the Pricing page and the plans: who prices low by owner type, whether price costs
# distribution, pro vs retail on the same brand, and how multi-brand owners ladder good / better / best.
GBB = lambda i: "g" if i <= LO else "best" if i >= HI else "b"
solid = [(cat, bid, v[1]) for cat, bs in pricing_out.items() for bid, v in bs.items() if v[0] >= 3]
learn = {"own": {}, "reach": [], "chan": None, "ladders": None}
by_own = collections.defaultdict(list)
for cat, bid, ix in solid:
    by_own[out_br[bid].get("oclass") or "unknown"].append(ix)
for k, v in by_own.items():
    if len(v) >= 5:
        learn["own"][k] = [len(v), r1(med(v), 3), r1(sum(i <= LO for i in v) / len(v), 3), r1(sum(i >= HI for i in v) / len(v), 3)]
pro_doors = lambda bid: sum(1 for d in out_br[bid].get("dists", []) if ctype(d) in ("pro", "mixed"))
for lo, hi, lbl in [(0, 2, "0-1"), (2, 5, "2-4"), (5, 10, "5-9"), (10, 10 ** 6, "10+")]:
    v = [ix for cat, bid, ix in solid if lo <= pro_doors(bid) < hi]
    if v:
        learn["reach"].append([lbl, len(v), r1(med(v), 3)])
gap = [v[5] / v[6] for bs in pricing_out.values() for v in bs.values() if v[5] and v[6]]
if gap:
    learn["chan"] = [len(gap), r1(med(gap), 3), r1(sum(g > 1.05 for g in gap) / len(gap), 3), r1(sum(g < 0.95 for g in gap) / len(gap), 3)]
lad = collections.defaultdict(lambda: collections.defaultdict(set))
for cat, bs in pricing_out.items():
    for bid, v in bs.items():
        if out_br[bid].get("owner_id"):
            lad[out_br[bid]["owner_id"]][cat].add((bid, GBB(v[1])))
multi = [v for cs in lad.values() for v in cs.values() if len(v) >= 2]
learn["ladders"] = [len(multi), sum(len({t for _, t in v}) >= 2 for v in multi), sum(len({t for _, t in v}) == 3 for v in multi)]
learn["n"] = len(solid)

print("price ladder:", len(priced), "prices,", sum(len(v) for v in pricing_out.values()), "brand-category pairs,",
      sum(1 for v in pricing_out.values() if len(v) >= 4), "categories with 4+ brands")

# ---- the Competition tab: one row per normalized category
cats_out = [{k: v for k, v in {
    "seg": s(r.get("Segment")), "cat": s(r["Category (normalized)"]), "n": r1(n(r.get("# core brands")), 0),
    "mix": {"pe": r1(n(r.get("PE-backed")), 0), "family": r1(n(r.get("Founder / family")), 0), "esop": r1(n(r.get("ESOP / employee")), 0),
            "sub_private": r1(n(r.get("Sub. of private co")), 0), "public": r1(n(r.get("Public / sub. of public")), 0),
            "other": r1(n(r.get("Unknown / house brand / other")), 0)},
    "pct_pe": r1(n(r.get("% PE-backed")), 3), "pct_priv": r1(n(r.get("% private (PE+family+ESOP+private sub)")), 3),
    "npw": r1(n0(r.get("# core brands (pro-weighted)")), 1), "pct_pe_pw": r1(n0(r.get("% PE (pro-wtd)")), 3), "pct_priv_pw": r1(n0(r.get("% private (pro-wtd)")), 3),
    "mixpw": {k: v for k, v in {"pe": r1(n0(r.get("PE-backed (pro-wtd)")), 1), "family": r1(n0(r.get("Family + ESOP (pro-wtd)")), 1),
              "public": r1(n0(r.get("Public / public sub (pro-wtd)")), 1)}.items() if v is not None},
    "avg_ret": r1(n0(r.get("Avg retail share of core brands")), 3),
    "skus": r1(n(r.get("# SKUs (core brands)")), 0), "top": s(r.get("Top 10 brands by distributor reach in this category (owner type)")),
    "pe_players": s(r.get("PE-backed players")), "fam_players": s(r.get("Family / ESOP players (roll-up candidates)")), "notes": s(r.get("Notes")),
    "npriced": r1(n(r.get("# priced items")), 0), "cpo": r1(n(r.get("Avg $/oz (category)")), 3), "nrev": r1(n(r.get("# brands with reviews")), 0),
}.items() if v not in ("", None)} for r in comp_rows]

# ---- competitors: brands selling the same kind of product, by TF-IDF over product categories and SKU names,
# nudged up when they sit on the same distributors' shelves. Brands under the same owner are left out.
NOISE_CAT = re.compile(r"directory|supplier|manufacturer list|not observed|catalog", re.I)
STOP = set("and or the of for with to in a an by on oz fl gal lb lbs ml l kg g ct pk pack case each ea in. x w per qt pt pint quart gallon can "
           "bottle tube cartridge aerosol spray black white red blue clear gray grey green yellow orange new size pc pcs".split())
def toks(text, own):
    w = [x[:-1] if len(x) > 4 and x.endswith("s") and not x.endswith("ss") else x for x in re.findall(r"[a-z][a-z\-]+", text.lower())]
    w = [x for x in w if x not in STOP and x not in own and len(x) > 2]
    return w + [a + " " + b for a, b in zip(w, w[1:])]
vec = {}
for bid, r in out_br.items():
    own = set(re.findall(r"[a-z]+", r["name"].lower()))
    tf = collections.Counter()
    for x in toks(r.get("cat", ""), own) + toks(r.get("hero", ""), own):
        tf[x] += 2
    for p in prod_by_brand.get(r["name"], []):
        c = s(p.get("Product category"))
        if c and not NOISE_CAT.search(c):
            for x in toks(c, own): tf[x] += 1
        for x in toks(s(p.get("Product name")), own): tf[x] += 0.5
    vec[bid] = tf
df = collections.Counter(x for tf in vec.values() for x in tf)
N = len(vec)
idf = {x: math.log(N / d) for x, d in df.items() if 1 < d < N * 0.2}
post = collections.defaultdict(list)
for bid, tf in vec.items():
    v = {x: math.log1p(c) * idf[x] for x, c in tf.items() if x in idf}
    norm_ = math.sqrt(sum(w * w for w in v.values())) or 1
    vec[bid] = {x: w / norm_ for x, w in v.items()}
    for x, w in vec[bid].items(): post[x].append((bid, w))
def owner_of(r):
    return r.get("owner_id") or key(r.get("owner") or "") or None
for bid, r in out_br.items():
    sim = collections.Counter()
    for x, w in vec[bid].items():
        for o, w2 in post[x]: sim[o] += w * w2
    mine, dset = owner_of(r), set(r.get("dists", []))
    scored = []
    for o, cs in sim.items():
        if o == bid or cs < 0.12 or (mine and owner_of(out_br[o]) == mine):
            continue
        od = set(out_br[o].get("dists", []))
        jac = len(dset & od) / len(dset | od) if dset | od else 0
        scored.append((cs * (1 + 0.6 * jac), o))
    scored.sort(reverse=True)
    if scored:
        r["rv"] = [o for _, o in scored[:10]]
print("competitors: brands with rivals", sum(1 for r in out_br.values() if r.get("rv")), "of", len(out_br))

# ---- judgments (Claude's read on the tests the data can't answer), merged when present
JF = os.path.join(os.path.dirname(os.path.abspath(__file__)), "judgments.json")
judgments = json.load(open(JF)) if os.path.exists(JF) else {}
for cid, j in judgments.items():
    if cid in out_co:
        out_co[cid]["judge"] = j

dist_out = [{"name": k, "market": market(k), "end_market": s(v.get("Primary end market")), "channel": channel(k), "ctype": ctype(k),
             "status": s(v.get("Scrape status")), "skus": int(n(v.get("SKUs scraped")) or 0),
             **{k2: v2 for k2, v2 in {"web": s(v.get("Website")), "secondary": s(v.get("Secondary markets"), 200),
                                      "scope": s(v.get("Chemical categories to scrape"), 200), "priority": s(v.get("Scrape priority")),
                                      "found": int(n(v.get("Brands found (pipeline)")) or 0), "method": s(v.get("Scrape method / limits"), 240)}.items() if v2}}
            for k, v in dists.items()]

# ---- write: shards sized to stay well under the 256 KB document cap
for d in ("data", "judge", "sku"):
    os.makedirs(f"{OUT}/{d}", exist_ok=True)
    for f in glob.glob(f"{OUT}/{d}/*.json"):
        os.remove(f)


def shard(items, prefix, field, limit=180_000, folder="data"):
    docs, cur, size = [], {}, 0
    for k, v in items.items():
        z = len(json.dumps(v)) + len(k) + 6
        if cur and size + z > limit:
            docs.append(cur); cur, size = {}, 0
        cur[k] = v; size += z
    if cur:
        docs.append(cur)
    for i, d in enumerate(docs):
        json.dump({field: d, "updated_at": NOW, "source": SEEDED}, open(f"{OUT}/{folder}/{prefix}-{i}.json", "w"), separators=(",", ":"))
    return {k: i for i, d in enumerate(docs) for k in d}


nco = len(set(shard(out_co, "co", "companies").values()))
ndt = len(set(shard(out_dt, "dt", "details").values()))
where = shard({k: v for k, v in out_sku.items() if v}, "p", "skus", folder="sku")  # SKUs load on demand, per brand
for bid, rec in out_br.items():
    if bid in where:
        rec["p"] = where[bid]
# ---- compact the brand records: every brand loads at start-up, so repeated strings become indexes
CATS = [c["cat"] for c in cats_out] + [c for c in sorted({r.get("ncat") for r in out_br.values()} | {k for r in out_br.values() for k in r.get("cats", {})})
                                       if c and c not in {x["cat"] for x in cats_out}]
CAT_IX = {c: i for i, c in enumerate(CATS)}
BASIS = {"Verified": "v", "Triage guess (unverified)": "t", "Prior knowledge (unverified)": "p"}
for r in out_br.values():
    r["nc"] = CAT_IX[r.pop("ncat")]
    r.pop("seg", None)
    if r.get("cats"):
        r["cats"] = {CAT_IX[k]: v for k, v in r["cats"].items()}
    r["wc"] = [e[2] if len(e) > 2 else e[:2] for e in r.get("wc", [])]  # linked rivals by id; the rest by name and label
    if not r["wc"]:
        r.pop("wc")
    wcids = {e for e in r.get("wc", []) if isinstance(e, str)}
    r["rv"] = [x for x in r.get("rv", []) if x not in wcids][:8]  # the shelf match adds to the market map's list
    r.pop("mkts", None)  # rebuilt in the app from each distributor's end market
    r.pop("tri_type", None)
    if same_owner(r.get("tri_owner"), r.get("owner")):
        r.pop("tri_owner", None)
    if not r["rv"]:
        r.pop("rv")
    b = r.pop("wb_basis", "")
    w, wt = r.pop("wb_owner", ""), r.pop("wb_otype", "")
    if b and b != "Verified" and w and not same_owner(r.get("owner"), w):
        r["wb"] = [w, wt, BASIS.get(b, b)]  # the workbook's best-available owner, when it says something different
for c in cats_out:
    c["i"] = CAT_IX[c["cat"]]
print("brand data KB", len(json.dumps(out_br)) // 1024)
qwhere = shard(dict(out_q), "q", "prices", folder="sku")  # prices load with the brand page, like SKUs
for bid, i in qwhere.items():
    out_br[bid]["q"] = i
nbr = len(set(shard(out_br, "br", "brands").values()))
json.dump({"cats": {CAT_IX[c]: {"b": v, **pricing_meta[c]} for c, v in pricing_out.items() if c in CAT_IX}, "learn": learn, "updated_at": NOW},
          open(f"{OUT}/data/pricing.json", "w"), separators=(",", ":"))
if margin_doc:
    json.dump({**margin_doc, "updated_at": NOW}, open(f"{OUT}/data/margin.json", "w"), separators=(",", ":"))
# ---- v29+: category market studies. The Categories tab scores every category on seven attractiveness criteria
# (weights in its header row) and sizes it; each category's own tab is a ten-section market study. Studies go to the
# "study" collection (one doc per category index, loaded when a category page opens); the scores and sizes ride on
# data/cats so the list can rank by them.
def bullets(rows):
    return [str(r[0]).lstrip("•? ").strip() for r in rows if isinstance(r[0], str) and r[0].strip()[:1] in ("•", "?")]
def norm_cat(x):
    return re.sub(r"[^a-z0-9]+", " ", str(x).lower().replace("'", "")).strip()
CAT_BY_NORM = {norm_cat(c): c for c in CATS}
cat_sum, study_out, crit = {}, {}, []
if "Categories" in wb.sheetnames:
    rows = list(wb["Categories"].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[1] == "Category")
    H = [str(x).strip() if x is not None else "" for x in rows[hi]]
    wrow = rows[hi - 1]
    crit = [[H[j], n0(wrow[j]) if wrow[j] is not None else 1] for j in range(4, 11)]
    for r in rows[hi + 1:]:
        c = CAT_BY_NORM.get(norm_cat(r[1])) if r[1] else None
        if not c:
            continue
        g = dict(zip(H, r))
        cat_sum[c] = {k: v for k, v in {"attr": r1(n0(g.get("Attractiveness (weighted)")), 2), "sc": [n0(r[j]) for j in range(4, 11)],
            "tam": r1(n0(g.get("US TAM point ($M)")), 0), "lo": r1(n0(g.get("TAM low ($M)")), 0), "hi": r1(n0(g.get("TAM high ($M)")), 0),
            "pro": r1(n0(g.get("Pro TAM ($M)")), 0), "cagr": r1(n0(g.get("CAGR %")), 2), "conf": s(g.get("TAM confidence")),
            "nb": r1(n0(g.get("# core brands (map)")), 0), "no": r1(n0(g.get("# owners (map)")), 0),
            "priv": r1(n0(g.get("% owners private (PE+family+ESOP+private sub)")), 3), "pe": r1(n0(g.get("% owners PE-backed")), 3),
            "ppo": r1(n0(g.get("Median list $/oz")), 3), "skus": r1(n0(g.get("# SKUs scraped")), 0)}.items() if v not in (None, "")}
for ws in wb.worksheets:
    rows = [r for r in ws.iter_rows(values_only=True)]
    if not any(isinstance(r[0], str) and r[0].startswith("1. WHAT IT IS") for r in rows if r):
        continue
    c = CAT_BY_NORM.get(norm_cat(rows[0][0]))
    if not c:
        print("study: no category for tab", ws.title); continue
    st = {"name": c}
    for i in (3, 5):  # header/value pairs under the title
        st.update({s(k): v for k, v in zip(rows[i], rows[i + 1]) if k and v not in (None, "")})
    secs, cur = {}, None
    for r in rows[7:]:
        c0 = r[0] if r else None
        if isinstance(c0, str) and re.match(r"^\d+\. ", c0):
            cur = int(c0.split(".")[0]); secs[cur] = []; continue
        if cur and r and any(x is not None for x in r):
            secs[cur].append(r)
    def subsplit(rs, labels):  # {label: rows} for sub-headed sections; rows before any label go under ""
        out, k = {"": []}, ""
        for r in rs:
            if isinstance(r[0], str) and r[0].strip() in labels and all(x is None for x in r[1:]):
                k = r[0].strip(); out[k] = []; continue
            out[k].append(r)
        return out
    t = lambda x, lim=1200: s(x, lim)
    one = subsplit(secs.get(1, []), {"Use occasions"})
    st["what"] = " ".join(t(r[0]) for r in one[""] if isinstance(r[0], str)); st["uses"] = bullets(one.get("Use occasions", []))
    two = subsplit(secs.get(2, []), {"Bull case", "Bear case / risks"})
    st["scores"] = [[t(r[0], 60), n0(r[1])] for r in two[""] if isinstance(r[0], str) and isinstance(r[1], (int, float)) and not str(r[0]).startswith("Weighted")]
    st["bull"], st["bear"] = bullets(two.get("Bull case", [])), bullets(two.get("Bear case / risks", []))
    em, ob, note = [], [], ""
    for r in secs.get(3, [])[1:]:
        if isinstance(r[0], str) and r[0].startswith("Research mix"):
            note = t(r[0], 600); continue
        if r[0] and isinstance(r[1], (int, float)):
            em.append([t(r[0], 60), r1(r[1], 3), t(r[2], 200)])
        if len(r) > 6 and r[5] and isinstance(r[6], (int, float)):
            ob.append([t(r[5], 60), r1(r[6], 3)])
    st["em"], st["obs"], st["em_note"] = em, ob, note
    four = subsplit(secs.get(4, []), {"How it was built"})
    st["size"] = {t(r[0], 60): r[1] for r in four[""] if r[0] and r[1] is not None}
    st["built"] = " ".join(t(r[0], 2000) for r in four.get("How it was built", []) if r[0])
    five = subsplit(secs.get(5, []), {"Users of these categories list this one as an adjacency:"})
    st["adj"] = [[t(r[0], 30), CAT_BY_NORM.get(norm_cat(r[1]), t(r[1], 80)), t(r[4] if len(r) > 4 else "", 160)] for r in five[""] if r[0] and r[1]]
    st["adj_in"] = [CAT_BY_NORM.get(norm_cat(x), t(x, 80)) for r in five.get("Users of these categories list this one as an adjacency:", []) for x in r if x]
    six = subsplit(secs.get(6, []), {"Private owners in this category (PE, family, ESOP, private subs) - ranked by brand reach"})
    otypes, tops = [], []
    for r in six[""][1:]:
        if r[0] and isinstance(r[1], (int, float)):
            otypes.append([t(r[0], 40), int(r[1])])
        if len(r) > 5 and r[3]:
            tops.append([t(r[3], 60), t(r[4], 30), t(r[5], 140)])
    priv = []
    for r in six.get("Private owners in this category (PE, family, ESOP, private subs) - ranked by brand reach", [])[1:]:
        if r[0]:
            priv.append([t(r[0], 100), t(r[1], 40), t(r[2], 60), r1(n0(r[3]), 1), r1(n0(r[4]), 0), r1(n0(r[5]), 0), r1(n0(r[6]), 3), r1(n0(r[7]), 3) if len(r) > 7 else None])
    st["otypes"], st["tops"], st["priv"] = otypes, tops, priv
    forms, pp = [], {}
    for r in secs.get(7, []):
        if r[0] and isinstance(r[1], (int, float)):
            forms.append([t(r[0], 40), r1(r[1], 3)])
        if len(r) > 5 and r[4] and isinstance(r[5], (int, float)):
            pp[t(r[4], 40)] = r1(r[5], 3)
    st["forms"], st["pp"] = forms, pp
    LBL8 = ["Key buying criteria", "Value chain & cost drivers", "Channels & specification", "Regulation", "Seasonality / cyclicality",
            "Substitutes / disruption", "Trends", "Consolidation & M&A", "Deals in this category (from Deals tab)"]
    eight = subsplit(secs.get(8, []), set(LBL8))
    st["know"] = [[k, bullets(eight[k]) or [t(" ".join(str(r[0]) for r in eight[k] if r[0]), 2500)]] for k in LBL8 if eight.get(k)]
    st["dq"] = bullets(secs.get(9, []))
    st["src"] = [[t(r[0], 160), t(r[3], 300) if len(r) > 3 else "", t(r[6], 300) if len(r) > 6 else ""] for r in secs.get(10, []) if r[0] and not str(r[0]).startswith("Profile researched")]
    st["foot"] = next((t(r[0], 600) for r in secs.get(10, []) if r[0] and str(r[0]).startswith("Profile researched")), "")
    # v31+: section 11 - chemistry, raw materials, cost structure, packaging bill of materials and secular score
    if secs.get(11):
        r11, mode, chem, rms, fbom, cost = secs[11], None, [], [], [], {}
        for j, r in enumerate(r11):
            c0 = s(r[0])
            if c0.startswith("Secular score"):
                st["secular"] = {"score": r1(n0(r[1]), 2), "rank": r1(n0(r[3]), 0), "wcagr": r1(n0(r[5]), 4)}; continue
            if c0 == "Key chemistry": mode = "chem"; continue
            if c0 == "Raw material": mode = "rm"; continue
            if c0.startswith("Cost structure"): mode = "cost"; continue
            if c0.startswith("Package-format mix"): mode = "fmt"; continue
            if c0 == "Format" and mode == "fmt": continue
            if mode == "chem" and c0: chem.append([c0[:200], s(r[1], 30), s(r[2], 300)])
            elif mode == "rm" and c0: rms.append([c0[:200], s(r[1], 40), r1(n0(r[2]), 0), s(r[3], 200), s(r[5] if len(r) > 5 else "", 200), s(r[7] if len(r) > 7 else "", 30)])
            elif mode == "cost":
                if c0 == "Raw materials" and r1 is not None:
                    vals = r11[j + 1]; cost = {k: r1(n0(v), 3) for k, v in zip(["rm", "pack", "labor", "freight", "gm"], vals[:5])}
                elif c0.startswith("Basis:"): st["cost_basis"] = c0[:700]
                elif c0.startswith("Shared inputs"): st["shared"] = re.sub(r"^Shared inputs / synergy:\s*", "", c0)[:1500]
                elif c0.startswith("Regulatory constraints"): st["regs"] = re.sub(r"^Regulatory constraints on inputs:\s*", "", c0)[:1500]
            elif mode == "fmt" and c0:
                if isinstance(r[1], (int, float)): fbom.append([c0[:40], r1(r[1], 3), s(r[2], 1500)])
                else: st["fmt_note"] = c0[:300]
        st.update({"chem": chem, "rms": rms, "cost": cost, "fbom": fbom})
    study_out[CAT_IX[c]] = st
    if c in cat_sum and em:
        cat_sum[c]["em"] = [[m, p] for m, p, _ in sorted(em, key=lambda x: -(x[1] or 0))[:3]]
os.makedirs(f"{OUT}/study", exist_ok=True)
for ci, st in study_out.items():
    json.dump({**st, "updated_at": NOW}, open(f"{OUT}/study/{ci}.json", "w"), separators=(",", ":"), default=str)
if study_out:
    print("category studies:", len(study_out), "| largest KB", max(len(json.dumps(v, default=str)) for v in study_out.values()) // 1024, "| scored", len(cat_sum))

# ---- v30+: end markets. The End Markets tab summarises the 14 end markets (each distributor belongs to one) and holds a
# category x end market matrix of scraped SKU shares; each "EM-" tab is a nine-section profile. Summary and matrix go
# to data/ems; profiles to the "emstudy" collection (doc id = end market index).
ems_doc, em_out = None, {}
if "End Markets" in wb.sheetnames:
    rows = list(wb["End Markets"].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "End market")
    H = [str(x).strip() if x is not None else "" for x in rows[hi]]
    summ = []
    for r in rows[hi + 1:]:
        if not r or not r[0]:
            break
        g = dict(zip(H, r))
        summ.append({k: v for k, v in {"name": s(r[0]), "nd": r1(n0(g.get("# distributors (map)")), 0), "nds": r1(n0(g.get("# with SKUs")), 0),
            "skus": r1(n0(g.get("Core SKUs")), 0), "nb": r1(n0(g.get("# brands")), 0), "priv": r1(n0(g.get("% SKUs private-owned")), 3),
            "fmt": s(g.get("Top format")), "pack": s(g.get("Top pack size")), "spend": r1(n0(g.get("Implied spend ($M)")), 0),
            "cats": s(g.get("Top categories (observed)"), 400), "dists": s(g.get("Top distributors (research)"), 400), "who": s(g.get("Who orders"), 900),
            "g": r1(n0(g.get("Growth score (1-5)")), 1), "d": r1(n0(g.get("Defensiveness (1-5)")), 1), "cagr": r1(n0(g.get("Base CAGR")), 4),
            "fast": s(g.get("Fastest-growing sub-segment"), 200)}.items() if v not in (None, "")})
    mi = next(i for i, r in enumerate(rows) if r and r[0] == "Category" and i > hi)
    MH = [s(x) for x in rows[mi]]
    matrix = {}
    for r in rows[mi + 1:]:
        c = CAT_BY_NORM.get(norm_cat(r[0])) if r and r[0] else None
        if c:
            matrix[CAT_IX[c]] = [r1(n0(r[1]), 0)] + [r1(n0(x), 3) for x in r[2:2 + len(summ)]]
    ems_doc = {"list": summ, "cols": MH[2:2 + len(summ)], "matrix": matrix, "note": s(rows[1][0], 900) if rows[1][0] else ""}
for ws in wb.worksheets:
    if not ws.title.startswith("EM-"):
        continue
    rows = [r for r in ws.iter_rows(values_only=True)]
    name = re.sub(r"^End market:\s*", "", s(rows[0][0]))
    ix = next((i for i, x in enumerate(ems_doc["list"]) if norm_cat(x["name"]) == norm_cat(name)), None) if ems_doc else None
    if ix is None:
        print("end market: no summary row for", ws.title); continue
    em = {"name": name, "labels": re.sub(r"^Raw distributor labels grouped here:\s*", "", s(rows[1][0], 400))}
    em.update({s(k): v for k, v in zip(rows[3], rows[4]) if k and v not in (None, "")})
    secs, cur = {}, None
    for r in rows[5:]:
        c0 = r[0] if r else None
        if isinstance(c0, str) and re.match(r"^\d+\. ", c0):
            cur = int(c0.split(".")[0]); secs[cur] = []; continue
        if cur and r and any(x is not None for x in r):
            secs[cur].append(r)
    def subs(rs, labels):
        out, k = {"": []}, ""
        for r in rs:
            if isinstance(r[0], str) and r[0].strip() in labels and all(x is None for x in r[1:]):
                k = r[0].strip(); out[k] = []; continue
            out[k].append(r)
        return out
    txt = lambda rs: " ".join(s(r[0], 3000) for r in rs if isinstance(r[0], str) and not r[0].strip().startswith("•"))
    one = subs(secs.get(1, []), {"Workforce / establishments", "Spend on these products", "Cyclicality", "Demand drivers"})
    em["who_are"] = txt(one[""]); em["workforce"] = txt(one.get("Workforce / establishments", [])); em["spend_txt"] = txt(one.get("Spend on these products", []))
    em["cyc"] = txt(one.get("Cyclicality", [])); em["drivers"] = bullets(one.get("Demand drivers", []))
    em["spend_note"] = " ".join(s(r[0], 1500) for r in one.get("Demand drivers", []) if isinstance(r[0], str) and r[0].startswith("Implied spend"))
    em["jobs"] = [[s(r[0], 160), s(r[1], 600)] for r in secs.get(2, [])[1:] if r[0] and r[1]]
    res, obs, mode = [], [], None
    for r in secs.get(3, []):
        c0 = s(r[0])
        if c0.startswith("Research ranking"): mode = "r"; continue
        if c0.startswith("Observed: category"): mode = "o"; continue
        if mode == "r" and c0:
            res.append([CAT_BY_NORM.get(norm_cat(c0), c0), s(r[1], 30), s(r[2], 300)])
        elif mode == "o" and c0:
            obs.append([CAT_BY_NORM.get(norm_cat(c0), c0), r1(n0(r[1]), 0), r1(n0(r[2]), 3), s(r[3], 30), s(r[4], 30), r1(n0(r[5]), 0), s(r[7] if len(r) > 7 else "", 300)])
    em["intro3"] = next((s(r[0], 600) for r in secs.get(3, []) if s(r[0]).startswith("Left:")), "")
    em["res"], em["obs"] = res, obs
    four = subs(secs.get(4, []), {"Ordering channels", "Order frequency", "Typical basket"})
    em["orders"] = txt(four[""]); em["channels"] = bullets(four.get("Ordering channels", [])); em["freq"] = txt(four.get("Order frequency", [])); em["basket"] = txt(four.get("Typical basket", []))
    rd, sd, mode = [], [], None
    for r in secs.get(5, []):
        c0 = s(r[0])
        if c0.startswith("Key distributors"): mode = "r"; continue
        if c0.startswith("Distributors in our scrape"): mode = "s"; continue
        if mode == "r" and c0:
            rd.append([c0[:160], s(r[1], 40), s(r[2], 400)])
        elif mode == "s" and c0:
            sd.append([c0[:160], s(r[1], 60), s(r[2], 30), s(r[3], 30), r1(n0(r[4]), 0), r1(n0(r[5]), 0), s(r[7] if len(r) > 7 else "", 120)])
    em["rdists"], em["sdists"] = rd, sd
    LB6 = {"Typical pack sizes (research)", "Price sensitivity", "Brand loyalty"}
    six, fm, pk, why, tb, mode, k6 = secs.get(6, []), [], [], [], [], "f", ""
    extra = {}
    for r in six:
        c0 = s(r[0])
        if c0.startswith("Observed format mix"): mode = "f"; continue
        if c0 in LB6 and all(x is None for x in r[1:]): mode, k6 = "x", c0; extra[k6] = []; continue
        if c0.startswith("Top brands here"): mode = "b"; continue
        if mode == "f":
            if c0 and isinstance(r[1], (int, float)): fm.append([c0[:40], r1(r[1], 3)])
            if len(r) > 4 and r[3] and isinstance(r[4], (int, float)): pk.append([s(r[3], 30), int(r[4])])
            if len(r) > 7 and r[7]: why.append(s(r[7], 300))
        elif mode == "x" and c0:
            extra[k6].append(c0)
        elif mode == "b" and c0:
            tb.append([c0[:80], r1(n0(r[1]), 0), r1(n0(r[2]), 0), CAT_BY_NORM.get(norm_cat(r[3]), s(r[3], 80)), s(r[4], 120), s(r[6] if len(r) > 6 else "", 40)])
    em["formats"], em["packs"], em["fwhy"], em["brands"] = fm, pk, why, tb
    em["packs_txt"] = " ".join(extra.get("Typical pack sizes (research)", [])); em["price_sens"] = " ".join(extra.get("Price sensitivity", [])); em["loyalty"] = " ".join(extra.get("Brand loyalty", []))
    ot, po = [], []
    for r in secs.get(7, [])[1:]:
        if r[0] and isinstance(r[1], (int, float)): ot.append([s(r[0], 40), int(r[1]), r1(n0(r[2]), 3)])
        if len(r) > 4 and r[4]: po.append([s(r[8] if len(r) > 8 and r[8] else r[4], 120), s(r[4], 200), r1(n0(r[7]), 3) if len(r) > 7 else None])
    em["otypes"], em["powners"] = ot, po
    LB8 = ["Regulation & standards", "Trends", "What wins here (implications for brand owners)", "Associations, shows, publications"]
    eight = subs(secs.get(8, []), set(LB8))
    em["know"] = [[k, bullets(eight[k]) or [txt(eight[k])]] for k in LB8 if eight.get(k)]
    em["src"] = [[s(r[0], 200), s(r[3], 300) if len(r) > 3 else "", s(r[6], 300) if len(r) > 6 else ""] for r in secs.get(9, []) if r[0] and not str(r[0]).startswith("Observed figures")]
    em["foot"] = next((s(r[0], 600) for r in secs.get(9, []) if r[0] and str(r[0]).startswith("Observed figures")), "")
    if secs.get(10):  # v31+: growth & cyclicality
        r10, gsubs, k10, lists = secs[10], [], None, {}
        for j, r in enumerate(r10):
            c0 = s(r[0])
            if c0.startswith("Growth score"):
                v = r10[j + 1]; em["growth"] = {"g": r1(n0(v[0]), 1), "d": r1(n0(v[2]), 1), "cagr": r1(n0(v[4]), 4)}; continue
            if c0 in ("Recurring maintenance vs new-project demand", "Downturn sensitivity"):
                em["recur" if c0.startswith("Recurring") else "downturn"] = s(r10[j + 1][0], 1500); continue
            if c0 == "Sub-segment": k10 = "subs"; continue
            if c0 in ("Secular tailwinds", "Secular headwinds", "Leading indicators to watch"): k10 = c0; lists[c0] = []; continue
            if c0.startswith("Growth sources"): em["gsrc"] = c0[:1500]; continue
            if k10 == "subs" and c0 and isinstance(r[1], (int, float)):
                gsubs.append([c0[:160], r1(r[1], 3), r1(n0(r[2]), 4), r1(n0(r[3]), 0), r1(n0(r[4]), 0), s(r[5], 400), s(r[7] if len(r) > 7 else "", 300)])
            elif k10 in lists and c0.startswith("•"):
                lists[k10].append(c0.lstrip("• ").strip())
        em["gsubs"], em["tail"], em["head"], em["lead"] = gsubs, lists.get("Secular tailwinds", []), lists.get("Secular headwinds", []), lists.get("Leading indicators to watch", [])
    em_out[ix] = em
if "Secular Screen" in wb.sheetnames:  # v31+: secular score per category (end-market growth, defensiveness, category growth)
    rows = list(wb["Secular Screen"].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "Category")
    H = [s(x) for x in rows[hi]]
    wi = next((i for i, r in enumerate(rows) if r and s(r[0]).startswith("Secular score weights")), None)
    for r in rows[hi + 1:]:
        if not r or not r[0]: break
        c = CAT_BY_NORM.get(norm_cat(r[0]))
        if c and c in cat_sum:
            g = dict(zip(H, r))
            cat_sum[c]["sec"] = {k: v for k, v in {"score": r1(n0(g.get("Secular score")), 2), "rank": r1(n0(g.get("Rank")), 0), "emg": r1(n0(g.get("EM growth (wtd)")), 2),
                "emd": r1(n0(g.get("EM defensiveness (wtd)")), 2), "emc": r1(n0(g.get("EM CAGR (wtd)")), 4), "cg": r1(n0(g.get("Category growth score")), 0)}.items() if v is not None}
    pi = next((i for i, r in enumerate(rows) if r and s(r[0]).startswith("HIGH-GROWTH POCKETS")), None)
    if ems_doc and pi is not None:
        pk = []
        for r in rows[pi + 2:]:
            if not r or not r[0]: continue
            pk.append([s(r[0], 60), s(r[1], 200), r1(n0(r[2]), 3), r1(n0(r[3]), 4), r1(n0(r[4]), 0), r1(n0(r[5]), 0), s(r[6], 400)])
        ems_doc["pockets"] = pk
    if ems_doc:
        ems_doc["sec_note"] = s(rows[1][0], 900)
        if wi is not None: ems_doc["sec_w"] = [n0(rows[wi][2]), n0(rows[wi][4]), n0(rows[wi][6])]
product_doc = None
if "Product Category" in wb.sheetnames:  # v31+: packaging, chemistry, raw materials and procurement synergies
    rows = list(wb["Product Category"].iter_rows(values_only=True))
    sec_at = {s(r[0])[:2]: i for i, r in enumerate(rows) if r and isinstance(r[0], str) and re.match(r"^[A-E]\. [A-Z]{3}", r[0])}
    def block(k, nxt):
        return rows[sec_at[k]:sec_at[nxt] if nxt in sec_at else len(rows)]
    fmts, cur = [], None
    for r in block("A.", "B.")[1:]:
        c0 = s(r[0])
        if not c0: continue
        if "  —  " in c0 and all(x is None for x in r[1:]):
            cur = {"name": c0.split("  —  ")[0].strip(), "full": c0.split("  —  ")[1].strip(), "desc": "", "comps": [], "kv": {}}; fmts.append(cur); continue
        if c0.startswith("Other formats") and isinstance(r[1], (int, float)):
            fmts.append({"name": "Other", "full": c0, "desc": "", "comps": [], "kv": {"Savings rate used in calculator": r[1]}}); cur = None; continue
        if cur is None: continue
        if c0 == "Component": continue
        if c0.startswith("Sources:"): cur["src"] = c0[:1500]; continue
        if c0 in ("Filling / co-packing", "Packaging % of COGS", "Synergy levers", "Typical savings (research)", "Savings rate used in calculator"):
            cur["kv"][c0] = r[1] if isinstance(r[1], (int, float)) else s(r[1], 1500); continue
        if not cur["comps"] and not cur["desc"] and all(x is None for x in r[1:]):
            cur["desc"] = c0[:1200]; continue
        cur["comps"].append([c0[:120]] + [s(x, 400) for x in r[1:9]])
    B = block("B.", "C."); bh = next(i for i, r in enumerate(B) if r and r[0] == "Category")
    BH = [s(x) for x in B[bh]]; fmix, fnote = {}, ""
    for r in B[bh + 1:]:
        c = CAT_BY_NORM.get(norm_cat(r[0])) if r and r[0] else None
        if c: fmix[CAT_IX[c]] = [s(r[1], 200), r1(n0(r[2]), 0)] + [r1(n0(x), 3) for x in r[3:len(BH)]]
        elif r and r[0]: fnote = s(r[0], 500)
    C = block("C.", "D."); ch = next(i for i, r in enumerate(C) if r and r[0] == "Category"); chem = {}
    for r in C[ch + 1:]:
        c = CAT_BY_NORM.get(norm_cat(r[0])) if r and r[0] else None
        if c: chem[CAT_IX[c]] = [s(x, 3000) if not isinstance(x, (int, float)) else x for x in r[1:10]]
    D = block("D.", "E."); dh = next(i for i, r in enumerate(D) if r and r[0] == "Raw-material family")
    rmf = [[s(r[0], 200), r1(n0(r[1]), 0), r1(n0(r[2]), 0), s(r[3], 600), [CAT_BY_NORM.get(norm_cat(x.strip()), x.strip()) for x in s(r[4], 4000).split(";") if x.strip()], s(r[5], 120)] for r in D[dh + 1:] if r and r[0]]
    E = block("E.", "Z."); eh = next(i for i, r in enumerate(E) if r and r[0] == "Category")
    rates = {s(r[0]): n0(r[1]) for r in E[1:eh] if r and r[0] and isinstance(r[1], (int, float))}
    syn = {}
    for r in E[eh + 1:]:
        c = CAT_BY_NORM.get(norm_cat(r[0])) if r and r[0] else None
        if c: syn[CAT_IX[c]] = [r1(n0(x), 4) for x in r[1:7]]
    product_doc = {"title": s(rows[0][0], 300), "intro": s(rows[1][0], 1200), "formats": fmts, "fcols": BH[3:], "fmix": fmix, "fnote": fnote,
                   "chem_cols": [s(x) for x in C[ch][1:10]], "chem": chem, "rmf": rmf, "rmf_note": s(D[1][0], 900) if len(D) > 1 else "",
                   "syn_note": s(E[1][0], 1200), "rates": rates, "syn": syn}
    # map each category's raw materials (study section 11) onto the section-D families, so a roll-up can see which
    # inputs its members share. Families are matched by keywords, preferring families that list the category.
    FAM_KW = [["binder", "latex", "acrylic emulsion", "vae", "alkyd", "emulsion", "pva", "acrylic polymer", "acrylic resin"],
        ["hydrocarbon", "mineral spirits", "aliphatic", "aromatic", "naphtha", "isoparaffin", "heptane", "hexane", "toluene", "xylene", "kerosene", "distillate"],
        ["pigment", "tio2", "titanium dioxide", "carbon black", "dye", "colorant", "fluorescent", "iron oxide"],
        ["carbonate", "caco3", "talc", "clay", "silica", "gypsum", "cement", "filler", "pumice", "sand", "kaolin", "mica", "perlite", "diatom", "aggregate", "vermiculite", "ceramic", "boron nitride", "fluoride"],
        ["propellant", "lpg", "dme", "co2", "hfc", "hfo", "propane", "butane", "152a", "134a", "nitrogen", "dimethyl ether"],
        ["epoxy", "polyester", "vinyl ester", "hardener", "curative", "bisphenol", "resin", "glycidyl", "diluent"],
        ["base oil", "mineral oil", "pao", "synthetic oil", "naphthenic", "paraffinic", "poe", "polyol ester", "oil", "pag", "pve", "base fluid"],
        ["fragrance", "limonene", "terpene", "citrus", "pine oil", "perfume"],
        ["tin", "zinc", "copper", "silver", "aluminium", "aluminum", "metal", "nickel", "alloy", "brass", "bronze", "cuprous", "oxide"],
        ["pesticide", "insecticide", "herbicide", "active ingredient", "pyrethroid", "glyphosate", "reagent", "technical-grade", "dpd", "indicator"],
        ["silicone", "pdms", "siloxane", "silane", "silanol"],
        ["isopropyl", "ipa", "ethanol", "methanol", "alcohol"],
        ["polyurethane", "mdi", "polyol", "prepolymer", "isocyanate", "stp", "ms polymer", "urethane", "tdi"],
        ["quat", "quaternary", "adbac", "ddac", "disinfect", "hypochlorite", "bleach", "peroxide", "biocide", "antimicrobial", "chlorine", "fungicide", "algaecide"],
        ["surfactant", "emulsifier", "ethoxylate", "las", "sles", "betaine", "amine oxide", "soap", "nonionic", "anionic", "sulfonate", "amphoteric"],
        ["nonwoven", "fibre", "fiber", "cellulose", "spunlace", "spunbond", "substrate", "cloth", "fabric", "pulp"],
        ["thickener", "rheology", "fumed silica", "thixotrop", "lithium soap", "gum", "wax", "xanthan", "bentonite", "cellulosic", "polyacrylamide", "polyacrylate"],
        ["rubber", "sbr", "butyl", "pvc", "polyethylene", "thermoplastic", "neoprene", "polychloroprene", "styrene-butadiene", "abs", "eva", "polypropylene"],
        ["acetone", "mek", "thf", "glycol ether", "ester", "ketone", "oxygenated", "butyl glycol", "dpm", "pnb", "cyclohexanone"],
        ["acid", "caustic", "hydroxide", "phosphoric", "hcl", "hydrochloric", "citric", "sulfamic", "alkali", "soda ash", "silicate", "metasilicate", "koh", "naoh"],
        ["salt", "chloride", "nacl", "cacl2", "mgcl2", "nitrite", "urea", "acetate", "sodium chloride", "potassium chloride", "magnesium"],
        ["inhibitor", "chelant", "edta", "glda", "mgda", "drier", "stabiliser", "stabilizer", "preservative", "antioxidant", "catalyst", "crosslinker", "adhesion promoter", "builder", "gluconate", "uv absorber", "additive", "flux activator", "activator", "azole", "triazole", "sulfite", "neutralising amine", "neutralizing amine", "deha"],
        ["water"],
        ["glycol", "propylene glycol", "ethylene glycol", "glycerin", "glycerine"],
        ["asphalt", "bitumen", "tar"],
        ["plasticizer", "plasticiser", "dinp", "didp", "dotp", "phthalate"],
        ["packaging", "container", "cartridge", "can "],
        ["cyanoacrylate", "methacrylate", "mma", "monomer", "acrylate"],
        ["ptfe", "fluoro", "pfpe", "fluoropolymer", "teflon"],
        ["enzyme", "microbial", "bacteria", "culture", "spore"],
        ["graphite", "carbon", "molybdenum", "moly"]]
    fam_names = [r[0] for r in rmf]
    kw_of = {i: FAM_KW[i] for i in range(min(len(FAM_KW), len(fam_names)))}
    cats_of = {i: set(r[4]) for i, r in enumerate(rmf)}
    def fam_for(name, cat):
        t = " " + name.lower().replace("/", " ").replace(",", " ") + " "
        sc = {i: sum((2 if " " in k.strip() else 1) for k in kws if k in t) for i, kws in kw_of.items()}
        sc = {i: v for i, v in sc.items() if v}
        if not sc: return None
        first = {i: min(t.find(k) for k in kw_of[i] if k in t) for i in sc}  # a mixed row goes to the material named first
        lead = min(first, key=first.get)
        pref = {i: v + (1 if cat in cats_of.get(i, ()) else 0) + (1.5 if i == lead else 0) for i, v in sc.items()}
        return max(pref, key=lambda i: (pref[i], -i))
    rmc, unm = {}, 0
    for ci, st in study_out.items():
        rows_ = []
        for rm in st.get("rms", []):
            f = fam_for(rm[0], st["name"])
            if f is None: unm += 1
            rows_.append([f, rm[2] or 0, rm[5] or "", rm[0][:120], rm[3][:200]])
        if rows_: rmc[ci] = rows_
    print("raw materials mapped to families:", sum(len(v) for v in rmc.values()), "rows,", unm, "unmatched")
    os.makedirs(f"{OUT}/product", exist_ok=True)  # the family mapping rides in its own doc (docs cap at 256 KB)
    json.dump({"fams": fam_names, "rmc": rmc, "updated_at": NOW}, open(f"{OUT}/product/rmc.json", "w"), separators=(",", ":"), default=str)
    json.dump({**product_doc, "updated_at": NOW}, open(f"{OUT}/product/main.json", "w"), separators=(",", ":"), default=str)
    print("product category:", len(fmts), "formats,", len(fmix), "format mixes,", len(chem), "chemistry rows,", len(rmf), "raw-material families,", len(syn), "synergy rows | KB", len(json.dumps(product_doc, default=str)) // 1024)
if ems_doc:
    os.makedirs(f"{OUT}/emstudy", exist_ok=True)
    for ix, em in em_out.items():
        json.dump({**em, "updated_at": NOW}, open(f"{OUT}/emstudy/{ix}.json", "w"), separators=(",", ":"), default=str)
    json.dump({**ems_doc, "updated_at": NOW}, open(f"{OUT}/data/ems.json", "w"), separators=(",", ":"), default=str)
    print("end markets:", len(ems_doc["list"]), "| profiles", len(em_out), "| matrix rows", len(ems_doc["matrix"]), "| largest KB", max(len(json.dumps(v, default=str)) for v in em_out.values()) // 1024)

json.dump({"cats": cats_out, "names": CATS, "study": {CAT_IX[c]: v for c, v in cat_sum.items()}, "crit": crit, "segs": {c: SEG_OF.get(c, "Unclassified") for c in CATS}, "updated_at": NOW}, open(f"{OUT}/data/cats.json", "w"), separators=(",", ":"))
json.dump({"deals": deals_out, "updated_at": NOW}, open(f"{OUT}/data/deals.json", "w"), separators=(",", ":"))
json.dump({"list": dist_out, "updated_at": NOW}, open(f"{OUT}/data/dists.json", "w"), separators=(",", ":"))
json.dump({"source": SEEDED, "updated_at": NOW, "co": nco, "dt": ndt, "br": nbr, "sku": len(set(where.values())), "q": len(set(qwhere.values()))}, open(f"{OUT}/data/index.json", "w"))
# the workbook's own IDs, so exports (the DNA workbook) can be joined back to the Master and Brands tabs
json.dump({"owners": {cid: oid for oid, cid in ids.items()}, "merged": {d: ids.get(c) for d, c in dupes.items()}, "brands": wb_brand_id},
          open(f"{OUT}/ids.json", "w"), indent=0)
for i in range(0, len(judge_in), 50):
    json.dump(judge_in[i:i + 50], open(f"{OUT}/judge/batch-{i // 50}.json", "w"), indent=1)

print("duplicate owners merged:", ", ".join(f"{d} into {c}" for d, c in dupes.items()) or "none")
print("owners", len(out_co), "verified", sum(1 for o in owners.values() if o["verified"]), "| judged", sum(1 for c in out_co.values() if "judge" in c))
print("brands", len(brands), "linked", sum(len(o["brands"]) for o in owners.values()), "| deals", len(deals_out),
      collections.Counter(d["kind"] for d in deals_out))
print("brand owners:", collections.Counter(b["status"] for b in out_br.values()), "| sku docs", len(set(where.values())))
print("shards: co", nco, "dt", ndt, "br", nbr, "| sizes KB", sorted(os.path.getsize(f) // 1024 for f in glob.glob(f"{OUT}/data/*.json")))
