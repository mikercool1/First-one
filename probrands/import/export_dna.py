"""Write the DNA test results as an Excel workbook that can be pasted into the market map.

    python3 probrands/import/export_dna.py Professional_Brands_Market_Map_v9.xlsx dna.json Probrands_DNA_v9.xlsx

dna.json is the app's own scoring (score, band and every check with its reason and source), dumped from the page so
the workbook matches the screen exactly. This script re-runs the importer for the workbook's Owner and Brand IDs and
the brand-level data checks.
"""
import collections, json, os, runpy, sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SRC, DNA, DEST = sys.argv[1:4]
here = os.path.dirname(os.path.abspath(__file__))
sys.argv = [os.path.join(here, "build_import.py"), SRC]
g = runpy.run_path(sys.argv[0], run_name="importer")  # prints its usual summary
ids = json.load(open(os.path.join(here, "out", "ids.json")))
dna = json.load(open(DNA))
out_br, prod_by_brand, channel, handheld, consumable = (g[k] for k in ("out_br", "prod_by_brand", "channel", "handheld", "consumable"))

TESTS = dna["tests"]
RES = ["Fail", "Partly", "Pass"]
BAND = {"g": "Green", "a": "Amber", "r": "Red"}
SRC_LABEL = {"data": "Data rule", "claude": "Claude's judgment", "you": "Your call"}
OWN = {"pe": "PE", "family": "Family", "esop": "ESOP", "private": "Private", "public": "Corporate", "sub_private": "Corporate", "unknown": "Unknown"}
STATUS = {"verified": "Owner verified", "prior": "Workbook prior (unverified)", "guess": "Claude's guess", "unknown": "Not identified"}
TRI = {"actionable": "Actionable candidate", "unknown": "Unknown brand", "sub_known": "Subsidiary, owner known", "sub_public": "Subsidiary of public co",
       "sub_large": "Subsidiary of other large co", "foreign": "Foreign maker", "house": "House brand", "consumer": "Consumer only",
       "out": "Out of scope", "not_chem": "Not a chemical"}
co = {c["id"]: c for c in dna["cos"]}
merged = collections.defaultdict(list)
for d, cid in ids["merged"].items():
    if cid:
        merged[cid].append(d)

HEAD = PatternFill("solid", fgColor="1B1C19")
FILL = {"Green": "C6EFCE", "Amber": "FFE7B3", "Red": "F8CBAD", "Pass": "C6EFCE", "Partly": "FFE7B3", "Fail": "F8CBAD"}


def sheet(ws, header, rows, widths, wrap=()):
    ws.append(header)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="F2B705"), HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 45
    for col in wrap:
        for c in ws[get_column_letter(col)][1:]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(row=2, column=3)
    ws.auto_filter.ref = ws.dimensions
    rng = f"A2:{get_column_letter(len(header))}{len(rows) + 1}"
    for word, color in FILL.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{word}"'], fill=PatternFill("solid", fgColor=color)))


wb = Workbook()

# ---- one row per owner: the DNA test as scored in the app
ws = wb.active
ws.title = "DNA by owner"
header = ["Rank", "Owner ID", "Also Master IDs (merged duplicates)", "Owner", "Ownership", "Owner type", "PE sponsor", "Parent", "Business model",
          "Primary trade", "HQ", "# brands", "# SKUs scraped", "DNA score (0-100)", "DNA band", "Out of scope", "Est. pro brands % of revenue"]
for t in TESTS:
    header += [f"{t['t']} ({t['w']} pts)", "Points", "Source", "Why"]
rows = []
ranked = sorted(dna["cos"], key=lambda c: (-c["score"], c["name"].lower()))
for i, c in enumerate(ranked, 1):
    oid = ids["owners"].get(c["id"], "")
    r = [i, oid if not oid.startswith("prior:") else "Not in Master", ", ".join(merged.get(c["id"], [])), c["name"], OWN.get(c["oclass"], "Unknown"),
         c["otype"], c["sponsor"], c["parent"], c["model"], c["trade"], c["hq"], c["nbrands"], c["skus"], c["score"], BAND[c["band"]],
         "Y" if c["out"] else "N", c["pro_pct"]]
    tw = {t["k"]: t["w"] for t in TESTS}
    for t in c["tests"]:
        r += [RES[t["v"]], tw[t["k"]] * t["v"] / 2, SRC_LABEL.get(t["src"], t["src"]), t["why"]]
    rows.append(r)
widths = [6, 9, 12, 34, 11, 18, 22, 24, 22, 20, 16, 8, 8, 9, 9, 8, 10] + [10, 7, 13, 48] * len(TESTS)
sheet(ws, header, rows, widths, wrap=[17 + 4 * j + 4 for j in range(len(TESTS))])

# ---- one row per brand: its owner's DNA plus the checks the scrape can answer for the brand itself
ws = wb.create_sheet("DNA by brand")
header = ["Brand ID", "Brand", "Owner ID", "Owner", "Owner status", "Ownership", "Owner DNA score", "Owner DNA band", "Owner out of scope",
          "Segment", "Category (normalized)", "SKUs scraped", "Handheld share", "Fits in your hand (brand)", "Pro distributors", "Pro-channel SKU share",
          "Sold through pro distributors (brand)", "Consumable share", "Used up and re-bought (brand)", "Brand data points (of 40)",
          "Triage bucket", "Top competitors (ownership)"]
rows = []
for bid, b in out_br.items():
    ps = prod_by_brand.get(b["name"], [])
    n = len(ps)
    hand = sum(handheld(p) for p in ps) / n if n else None
    cons = sum(consumable(p) for p in ps) / n if n else None
    pro = sum(channel(g["s"](p.get("Distributor"))) != "retail" for p in ps) / n if n else None
    ndp = sum(1 for d in b.get("dists", []) if channel(d) != "retail")
    vh = None if hand is None else 2 if hand >= 0.75 else 1 if hand >= 0.45 else 0
    vd = None if pro is None else 2 if ndp >= 4 and pro >= 0.7 else 1 if ndp >= 2 and pro >= 0.5 else 0
    vc = None if cons is None else 2 if cons >= 0.8 else 1 if cons >= 0.5 else 0
    pts = None if n == 0 else 15 * vh / 2 + 15 * vd / 2 + 10 * vc / 2
    c = co.get(b.get("owner_id")) if b.get("owner_id") else None
    oid = ids["owners"].get(b.get("owner_id"), "") if c else ""
    status = "Workbook triage guess" if b.get("src") == "triage" else STATUS[b["status"]]
    rv = "; ".join(f"{out_br[x]['name']} ({OWN.get(out_br[x].get('oclass'), 'Unknown')})" for x in b.get("rv", [])[:5])
    rows.append([ids["brands"].get(bid, ""), b["name"], oid if not oid.startswith("prior:") else "Not in Master", b.get("owner", ""), status,
                 OWN.get((c or b).get("oclass"), "Unknown"), c["score"] if c else None, BAND[c["band"]] if c else "Not rated",
                 ("Y" if c["out"] else "N") if c else None, g["SEG_OF"].get(g["CATS"][b["nc"]], "Unclassified") if "nc" in b else "", g["CATS"][b["nc"]] if "nc" in b else b.get("cat", ""), n,
                 None if hand is None else round(hand, 3), RES[vh] if vh is not None else "", ndp, None if pro is None else round(pro, 3),
                 RES[vd] if vd is not None else "", None if cons is None else round(cons, 3), RES[vc] if vc is not None else "", pts,
                 TRI.get(b.get("tri"), ""), rv])
rows.sort(key=lambda r: (-(r[6] if r[6] is not None else -1), r[1].lower()))
sheet(ws, header, rows, [9, 26, 9, 30, 20, 11, 9, 9, 8, 18, 26, 8, 9, 10, 9, 9, 11, 9, 10, 9, 14, 60])
for col in "MPR":
    for cell in ws[col][1:]:
        cell.number_format = "0%"

# ---- how it's scored
ws = wb.create_sheet("How it's scored")
ws.append(["Check", "Points", "Question", "How it is scored"])
how = {"hand": "Data rule: share of scraped SKUs that are aerosols, tubes, cartridges, wipes, markers, kits, or bottles/cans/powders up to 1 gal, 4 L or 10 lb. Pass 75%+, Partly 45%+.",
       "dist": "Data rule: pro (non-retail) distributors carrying it and share of SKUs in pro channels. Pass 4+ distributors and 70%+, Partly 2+ and 50%+.",
       "pull": "Claude's judgment (data fallback: widest brand on 6+ pro distributors = Pass, 3+ = Partly).",
       "itw": "Claude's judgment: handheld, formulated, consumable, MRO/trade channel.",
       "major": "Claude's judgment. Fail means pro brands are a small share of the owner's business and the owner is out of scope.",
       "cons": "Data rule: share of SKUs that are consumables, not tools, dispensers or equipment. Pass 80%+, Partly 50%+.",
       "own": "Data rule from the business model: private label, contract fill or distributor-owned house brand fails."}
for t in TESTS:
    ws.append([t["t"], t["w"], t["q"], how[t["k"]]])
ws.append([])
for line in ["Each check scores Pass (full points), Partly (half) or Fail (0). DNA score = sum of points, 0-100.",
             "Bands: Green 70+, Amber 45-69, Red under 45.",
             "Precedence: your call in the app, then a Claude re-check from the app, then the imported Claude judgment, then the data rule.",
             "The DNA test is scored per owner. 'DNA by brand' gives each brand its owner's score, plus the three data checks worked out for the brand alone (40 of the 100 points).",
             "Brands with no profiled owner (Claude's guess of an owner outside the map, or not identified) are 'Not rated' at owner level.",
             "Join keys: Owner ID matches the Master tab; Brand ID matches the Brands tab. 'Not in Master' owners come from the Brands tab's unverified prior.",
             f"Source: {g['SEEDED']}."]:
    ws.append([line])
for i, w in enumerate([30, 8, 60, 90], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for c in ws[1]:
    c.font, c.fill = Font(bold=True, color="F2B705"), HEAD
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
wb.save(DEST)
print("wrote", DEST, "| owners", len(ranked), "| brands", len(rows))
